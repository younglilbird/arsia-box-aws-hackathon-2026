"""Validate downloadable artifacts, real link relationships and hostile text."""

import io
import unittest
import zipfile
from xml.etree import ElementTree

from docx import Document
from pypdf import PdfReader

from app.exports import make_export


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.url = "https://app.box.com/file/12345?version=2&view=review"
        self.run = {
            "employee_name": "Alex Chen",
            "year": 2026,
            "language": "en",
            "report_markdown": "## Annual overview\n\nI owned **instrumentation** and verification. The team observed median task time fall from 10 to 8 minutes in its pilot.[E1]\n\n### Growth\n\n- Collaborated with support agents.\n- Personal reflection requires my input.",
            "evidence": [{"id": "E1", "source_id": "result", "source_name": "Atlas pilot results", "claim": "20% actual improvement", "quote": "300 participants; median 10 to 8 minutes."}],
            "sources": [{"id": "result", "name": "Atlas results", "url": self.url}],
            "reviewer_confirmed": False,
        }

    def test_editable_docx_has_link_and_review_state(self):
        data, name, mime = make_export(self.run, "docx")
        document = Document(io.BytesIO(data))
        self.assertTrue(name.endswith(".docx"))
        self.assertIn("wordprocessingml", mime)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("Draft - not yet confirmed by employee", text)
        self.assertIn("instrumentation", text)
        self.assertIn("Evidence and sources", text)
        targets = [rel.target_ref for rel in document.part.rels.values() if rel.reltype.endswith("/hyperlink")]
        self.assertEqual(targets, [self.url])
        document.paragraphs[3].add_run(" Editable after download.")
        buffer = io.BytesIO()
        document.save(buffer)
        self.assertGreater(len(buffer.getvalue()), 1000)

    def test_pdf_english_text_and_clickable_sources(self):
        data, name, mime = make_export(self.run, "pdf")
        self.assertTrue(data.startswith(b"%PDF"))
        self.assertTrue(name.endswith(".pdf"))
        self.assertEqual(mime, "application/pdf")
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(page.extract_text() for page in reader.pages)
        self.assertIn("Annual self-review", text)
        self.assertIn("instrumentation", text)
        self.assertEqual(len(reader.pages), 1)  # short body does not force another source page
        fonts = [item.get_object() for item in reader.pages[0]["/Resources"]["/Font"].values()]
        self.assertTrue(any(item.get("/BaseFont") == "/Helvetica" for item in fonts), "PDF uses a standard portable English font")
        links = []
        for page in reader.pages:
            for annotation in page.get("/Annots", []):
                action = annotation.get_object().get("/A", {})
                if action.get("/S") == "/URI":
                    links.append(action["/URI"])
        self.assertIn(self.url, links)

    def test_word_and_pdf_omit_only_the_redundant_leading_report_title(self):
        title = "Annual Self-Review for Alex Chen, 2026"
        self.run["report_markdown"] = "\n\n# " + title + "\n\n" + self.run["report_markdown"] + "\n\n# Next-year priorities\n\nConfirm my next-year goals."
        original = self.run["report_markdown"]
        document = Document(io.BytesIO(make_export(self.run, "docx")[0]))
        word_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        reader = PdfReader(io.BytesIO(make_export(self.run, "pdf")[0]))
        pdf_text = "\n".join(page.extract_text() for page in reader.pages)
        for text in (word_text, pdf_text):
            self.assertNotIn(title, text)
            self.assertEqual(text.count("Annual self-review"), 1)
            self.assertIn("Annual overview", text)
            self.assertIn("Next-year priorities", text)
            self.assertIn("instrumentation", text)
            self.assertIn("[1]", text)
        targets = [rel.target_ref for rel in document.part.rels.values() if rel.reltype.endswith("/hyperlink")]
        self.assertIn(self.url, targets)
        pdf_targets = [annotation.get_object().get("/A", {}).get("/URI") for page in reader.pages for annotation in page.get("/Annots", [])]
        self.assertIn(self.url, pdf_targets)
        self.assertEqual(self.run["report_markdown"], original)

    def test_distinct_leading_title_and_nonleading_titles_are_preserved(self):
        for body in ("# Draft review notes\n\nConfirm the claims. [E1]", "Opening context.\n\n# Annual self-review\n\nA later section. [E1]", "# Annual self-review challenges\n\nConfirm the challenges. [E1]"):
            with self.subTest(body=body):
                self.run["report_markdown"] = body
                expected = next(line[2:] for line in body.splitlines() if line.startswith("# "))
                document = Document(io.BytesIO(make_export(self.run, "docx")[0]))
                self.assertIn(expected, [paragraph.text for paragraph in document.paragraphs[3:]])
                reader = PdfReader(io.BytesIO(make_export(self.run, "pdf")[0]))
                text = "\n".join(page.extract_text() for page in reader.pages)
                self.assertIn(expected, text)

    def test_word_and_pdf_compactly_group_only_cited_sources(self):
        self.run["report_markdown"] += "\n\nThe same source supports a second point.[E2]"
        self.run["evidence"].append({"id": "E2", "source_id": "result", "source_name": "Atlas pilot results", "claim": "DUPLICATE LONG CLAIM", "quote": "DO NOT REPEAT FULL QUOTES " * 500})
        self.run["evidence"].append({"id": "E3", "source_id": "other", "source_name": "UNREFERENCED SOURCE", "quote": "UNREFERENCED QUOTE"})
        self.run["sources"].append({"id": "other", "name": "UNREFERENCED SOURCE", "url": "https://app.box.com/file/999"})
        document = Document(io.BytesIO(make_export(self.run, "docx")[0]))
        word_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertEqual(word_text.count("Atlas pilot results"), 1)
        self.assertNotIn("UNREFERENCED SOURCE", word_text)
        self.assertNotIn("DO NOT REPEAT", word_text)
        reader = PdfReader(io.BytesIO(make_export(self.run, "pdf")[0]))
        pdf_text = "\n".join(page.extract_text() for page in reader.pages)
        self.assertEqual(pdf_text.count("Atlas pilot results"), 1)
        self.assertIn("[1], [2]", pdf_text)
        self.assertNotIn("[E1]", pdf_text)
        self.assertNotIn("UNREFERENCED SOURCE", pdf_text)
        self.assertNotIn("DO NOT REPEAT", pdf_text)
        self.assertEqual(len(reader.pages), 1)

    def test_markdown_reference_links_and_confirmed_label(self):
        self.run["language"] = "en"
        self.run["reviewer_confirmed"] = True
        data, name, _ = make_export(self.run, "markdown")
        text = data.decode("utf-8")
        self.assertIn(f"[\\[1\\]](<{self.url}>)", text)
        self.assertNotIn("E1", text)
        self.assertIn("Confirmed by employee", text)
        self.assertNotIn("not yet confirmed", text)
        self.assertTrue(name.endswith(".md"))

    def test_reference_display_inside_bold_text_hides_internal_ids(self):
        self.run["report_markdown"] = "## Summary\n\n**The pilot supports this observation [E1].**"
        markdown = make_export(self.run, "md")[0].decode("utf-8")
        self.assertNotIn("E1", markdown)
        self.assertIn(f"[\\[1\\]](<{self.url}>)", markdown)
        reader = PdfReader(io.BytesIO(make_export(self.run, "pdf")[0]))
        text = "\n".join(page.extract_text() for page in reader.pages)
        self.assertIn("[1]", text)
        self.assertNotIn("E1", text)

    def test_model_markup_is_plain_text_and_unsafe_url_is_inert(self):
        payload = '<script>alert("x")</script> <img src=x onerror=boom> [click](javascript:alert(1)) & <b>text</b>\x00'
        self.run["report_markdown"] = payload + " [E1]"
        self.run["evidence"][0]["url"] = "javascript:alert(1)"
        self.run["sources"][0]["url"] = "data:text/html,boom"
        markdown = make_export(self.run, "md")[0].decode("utf-8")
        self.assertNotIn("<script>", markdown)
        self.assertIn("&lt;script&gt;", markdown)
        self.assertNotIn("](javascript:", markdown)
        self.assertNotIn("\x00", markdown)
        data = make_export(self.run, "docx")[0]
        document = Document(io.BytesIO(data))
        self.assertIn('<script>alert("x")</script>', document.paragraphs[3].text)
        targets = [rel.target_ref for rel in document.part.rels.values() if rel.reltype.endswith("/hyperlink")]
        self.assertEqual(targets, [])
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            ElementTree.fromstring(archive.read("word/document.xml"))
        pdf = make_export(self.run, "pdf")[0]
        self.assertIn("<script>", PdfReader(io.BytesIO(pdf)).pages[0].extract_text())

    def test_source_url_rejects_controls_and_user_info(self):
        for url in ("https://user:password@example.com/file", "https://example.com/a\nBAD", "file:///tmp/report", "https://example.com\\@evil.test/", 'https://example.com/><img/src=x>', 'https://example.com/"bad'):
            self.run["sources"][0]["url"] = url
            data = make_export(self.run, "docx")[0]
            document = Document(io.BytesIO(data))
            self.assertFalse(any(rel.reltype.endswith("/hyperlink") for rel in document.part.rels.values()))

    def test_empty_reports_and_unknown_formats_are_rejected(self):
        with self.assertRaises(ValueError):
            make_export({"report_markdown": " "}, "pdf")
        with self.assertRaises(ValueError):
            make_export(self.run, "html")


if __name__ == "__main__":
    unittest.main()
