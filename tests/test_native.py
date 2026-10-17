"""Exercise the native poll boundary without network or model charges."""

from copy import deepcopy
from io import BytesIO
import json
from types import SimpleNamespace

from docx import Document
from pypdf import PdfReader
import pytest

from app.native import NativeWorkflowError, _review_document, process_once


CONFIG = {
    "root_folder_id": "100", "category_folders": {"Projects": "101", "Peer Feedback": "102", "Annual Plans": "103"},
    "reports_folder_id": "104", "employee_name": "Alex Chen", "year": 2026, "language": "en",
}
PRIVATE = "PRIVATE SOURCE TEXT THAT MUST NEVER ENTER THE STATE STORE"


def file(identifier="201", name="Project.docx", version="1", **extra):
    return {"id": identifier, "type": "file", "name": name, "size": 100,
            "file_version": {"id": version}, **extra}


def folder(identifier, name):
    return {"id": identifier, "type": "folder", "name": name}


class Store:
    def __init__(self, initial=None):
        self.value = deepcopy(initial or {})
        self.history = []

    def read(self):
        return deepcopy(self.value)

    def write(self, value):
        self.value = deepcopy(value)
        self.history.append(deepcopy(value))


class Clock:
    def __init__(self):
        self.now = 1000.0

    def __call__(self):
        return self.now


class Box:
    def __init__(self):
        self.entries = {"101": [file()], "102": [], "103": [], "104": [file("999", "Generated.pdf")]}
        self.listed, self.descriptions, self.uploads, self.questions = [], [], [], []

    def list_folder(self, identifier):
        self.listed.append(identifier)
        return {"id": identifier, "entries": deepcopy(self.entries.get(identifier, [])), "truncated": False}

    def request(self, method, path, **kwargs):
        assert method == "PUT"
        assert path == "/2.0/folders/100"
        self.descriptions.append(kwargs["json"]["description"])
        return {"id": "100"}

    def read_text(self, selected):
        return PRIVATE

    def ask(self, identifiers, prompt):
        self.questions.append((identifiers, prompt))
        return {"answer": "Checked", "citations": []}

    def upload(self, folder_id, name, data, mime):
        assert folder_id == "104"
        self.uploads.append({"folder_id": folder_id, "name": name, "data": data, "mime": mime})
        return {"id": str(800 + len(self.uploads)), "name": name}


class Workflow:
    def __init__(self, failure=None):
        self.calls, self.failure, self.requests, self.discoveries = 0, failure, [], []

    def __call__(self, box, settings, request, update):
        self.calls += 1
        self.requests.append(request)
        if self.failure:
            raise self.failure
        discovery = box.discover(request["folder_id"], request["max_files"])
        self.discoveries.append(discovery)
        sources = [{**item, "text": box.read_text(item)} for item in discovery["files"]]
        box.ask([item["id"] for item in sources], "Check these sources")
        update(stage="extract", progress=25, message=PRIVATE, sources=sources)
        evidence = [{"id": "E1", "source_id": sources[0]["id"], "source_name": sources[0]["name"],
                     "quote": PRIVATE, "claim": "Alex implemented the measurement plan."}]
        return {
            "report_markdown": "# Annual review\n\nI implemented the measurement plan. [E1]",
            "sources": sources, "evidence": evidence,
            "review": {"issues": [{"severity": "review", "message": "Confirm personal ownership.", "evidence_ids": ["E1"]}],
                       "revision_performed": True,
                       "box_ai": {"status": "completed", "answer": "Expected impact should stay separate from measured impact.",
                                  "file_ids": [sources[0]["id"]], "citations": [{"source": sources[0]["id"]}],
                                  "draft_truncated": True, "note": "Box AI saw the first 8,000 draft characters."}},
            "missing_information": ["Confirm the reporting period."],
            "coverage": {"selected_files": len(sources), "analyzed_files": len(sources), "warnings": [PRIVATE], "skipped": discovery["skipped"]},
            "usage": {"input_tokens": 100},
        }


def poll(box, store, clock, workflow, config=None):
    return process_once(box, SimpleNamespace(box_client_secret="NEVER LOG THIS SECRET"), config or CONFIG,
                        store, workflow=workflow, clock=clock)


def settle_and_run(box, store, clock, workflow):
    assert poll(box, store, clock, workflow)["status"] == "settling"
    clock.now += 20
    return poll(box, store, clock, workflow)


def test_quiet_period_resets_on_each_material_change_and_completed_inputs_are_idempotent():
    box, store, clock, workflow = Box(), Store(), Clock(), Workflow()
    initial = poll(box, store, clock, workflow)
    assert initial["status"] == "settling" and workflow.calls == 0
    clock.now += 19
    assert poll(box, store, clock, workflow)["status"] == "settling"
    box.entries["101"][0]["file_version"]["id"] = "2"
    clock.now += 1
    changed = poll(box, store, clock, workflow)
    assert changed["last_changed_at"] == 1020
    clock.now += 19
    assert poll(box, store, clock, workflow)["status"] == "settling"
    assert workflow.calls == 0
    clock.now += 1
    result = poll(box, store, clock, workflow)
    assert result["status"] == "completed"
    assert workflow.calls == 1
    clock.now += 60
    assert poll(box, store, clock, workflow)["status"] == "unchanged"
    assert store.value["status"] == "completed"
    assert workflow.calls == 1


def test_real_draft_exports_are_clean_and_keep_private_review_state():
    box, store, clock, workflow = Box(), Store(), Clock(), Workflow()
    result = settle_and_run(box, store, clock, workflow)
    assert result["reviewer_confirmed"] is False
    assert result["report_ids"] == {"docx": "801", "pdf": "802", "review_notes": "803"}
    assert all(item["name"].startswith("Annual-Review-Alex-Chen-2026-Draft-") for item in box.uploads)
    word = Document(BytesIO(box.uploads[0]["data"]))
    assert "Draft - not yet confirmed by employee" in "\n".join(p.text for p in word.paragraphs)
    assert not [rel for rel in word.part.rels.values() if rel.reltype.endswith("/hyperlink")]
    assert "[1]" not in "\n".join(p.text for p in word.paragraphs)
    assert word.core_properties.author == "Alex Chen"
    pdf = PdfReader(BytesIO(box.uploads[1]["data"]))
    assert "not yet confirmed" in "\n".join(page.extract_text() for page in pdf.pages)
    notes = Document(BytesIO(box.uploads[2]["data"]))
    note_text = "\n".join(p.text for p in notes.paragraphs)
    assert "Confirm personal ownership" in note_text and PRIVATE in note_text
    assert "Box AI source review" in note_text
    assert "Status: completed. Source files checked: 1. Returned citations: 1." in note_text
    assert "Expected impact should stay separate from measured impact." in note_text
    assert "Draft input was truncated for Box AI" in note_text
    serialized = json.dumps(store.history)
    assert PRIVATE not in serialized and "NEVER LOG THIS SECRET" not in serialized
    assert "report_markdown" not in serialized and "quote" not in serialized
    assert result["evidence_count"] == 1 and result["revision_performed"] is True
    assert all(description.startswith("Status: ") for description in box.descriptions)


def test_box_ai_unavailable_is_explicit_in_native_review_notes():
    run = {"report_markdown": "Draft", "review": {"box_ai": {"status": "unavailable", "message": "Box AI was unavailable; other checks continued."}}}
    note = _review_document(run)["report_markdown"]
    assert "Status: unavailable. Source files checked: 0. Returned citations: 0." in note
    assert "Box AI was unavailable; other checks continued." in note
    assert "do not certify" in note


def test_recursion_filters_inputs_propagates_categories_and_never_reads_reports():
    box, store, clock, workflow = Box(), Store(), Clock(), Workflow()
    box.entries["101"] = [folder("105", "Nested project"), folder("104", "Reports moved under Projects"), file("202", "Unsupported.txt")]
    box.entries["105"] = [file("203", "Results.PDF")]
    box.entries["102"] = [file("204", "Feedback.docx")]
    box.entries["103"] = [file("205", "Goals.pdf")]
    settle_and_run(box, store, clock, workflow)
    assert "104" not in box.listed
    discovered = workflow.discoveries[0]
    assert [(item["id"], item["document_category"]) for item in discovered["files"]] == [
        ("203", "Projects"), ("204", "Peer Feedback"), ("205", "Annual Plans")]
    assert any(item["id"] == "202" for item in discovered["skipped"])
    request = workflow.requests[0]
    assert request["folder_id"] == "100" and request["language"] == "en"
    assert request["max_files"] == 25 and request["notes"] == ""
    assert box.questions[0][0] == ["203", "204", "205"]


def test_empty_or_unsupported_sources_wait_without_workflow_or_exports():
    box, store, clock, workflow = Box(), Store(), Clock(), Workflow()
    box.entries["101"] = [file(name="Slides.pptx")]
    result = poll(box, store, clock, workflow)
    assert result["status"] == "waiting" and result["source_count"] == 0
    assert workflow.calls == 0 and not box.uploads
    assert "Word/PDF" in box.descriptions[-1]


def test_over_limit_files_are_transparent_and_still_part_of_change_detection():
    box, store, clock, workflow = Box(), Store(), Clock(), Workflow()
    box.entries["101"] = [file(str(200 + i), f"Project-{i}.docx") for i in range(27)]
    result = settle_and_run(box, store, clock, workflow)
    assert result["source_count"] == 27 and result["selected_count"] == 25
    assert len(workflow.discoveries[0]["files"]) == 25
    assert workflow.discoveries[0]["truncated"] is True
    assert len(workflow.discoveries[0]["skipped"]) == 2
    assert any("25" in warning for warning in result["warnings"])
    box.entries["101"][-1]["file_version"]["id"] = "2"
    clock.now += 60
    assert poll(box, store, clock, workflow)["status"] == "settling"


def test_failure_is_sanitized_and_same_fingerprint_is_not_retried():
    box, store, clock = Box(), Store(), Clock()
    workflow = Workflow(RuntimeError("provider secret: " + PRIVATE))
    assert poll(box, store, clock, workflow)["status"] == "settling"
    clock.now += 20
    with pytest.raises(NativeWorkflowError) as raised:
        poll(box, store, clock, workflow)
    assert PRIVATE not in str(raised.value)
    assert store.value["status"] == "failed"
    assert PRIVATE not in json.dumps(store.history)
    clock.now += 60
    assert poll(box, store, clock, workflow)["status"] == "failed"
    assert workflow.calls == 1
    box.entries["101"][0]["file_version"]["id"] = "2"
    assert poll(box, store, clock, workflow)["status"] == "settling"
    workflow.failure = None
    clock.now += 20
    assert poll(box, store, clock, workflow)["status"] == "completed"
    assert workflow.calls == 2


def test_new_source_version_during_generation_blocks_publication_and_next_version_runs():
    box, store, clock, workflow = Box(), Store(), Clock(), Workflow()
    def changing_workflow(adapter, settings, request, update):
        result = workflow(adapter, settings, request, update)
        box.entries["101"][0]["file_version"]["id"] = "2"
        return result
    first = poll(box, store, clock, changing_workflow)
    assert first["status"] == "settling"
    clock.now += 20
    with pytest.raises(NativeWorkflowError) as raised:
        poll(box, store, clock, changing_workflow)
    assert raised.value.code == "source_changed"
    assert "No draft was published" in str(raised.value)
    assert PRIVATE not in str(raised.value) and "201" not in str(raised.value)
    assert not box.uploads
    assert store.value["status"] == "failed"
    assert store.value["failed_fingerprint"] == first["fingerprint"]
    assert store.value["error"]["code"] == "source_changed"
    assert store.value["report_ids"] == {} and not store.value.get("completed_fingerprint")
    assert workflow.discoveries[0]["files"][0]["file_version"]["id"] == "1"
    assert "No draft was published" in box.descriptions[-1]
    clock.now += 60
    settling = poll(box, store, clock, workflow)
    assert settling["status"] == "settling"
    assert settling["fingerprint"] != first["fingerprint"]
    clock.now += 20
    completed = poll(box, store, clock, workflow)
    assert completed["status"] == "completed"
    assert completed["completed_fingerprint"] == settling["fingerprint"]
    assert workflow.discoveries[1]["files"][0]["file_version"]["id"] == "2"
    assert len(box.uploads) == 3
    assert workflow.calls == 2


def test_interrupted_processing_run_is_not_restarted_automatically():
    box, store, clock, workflow = Box(), Store(), Clock(), Workflow()
    first = poll(box, store, clock, workflow)
    store.value.update(status="processing", fingerprint=first["fingerprint"], run_id="interrupted-run")
    clock.now += 60
    assert poll(box, store, clock, workflow)["status"] == "processing"
    assert workflow.calls == 0


def test_partial_upload_is_retained_on_failure_without_a_costly_retry():
    box, store, clock, workflow = Box(), Store(), Clock(), Workflow()
    original_upload = box.upload
    def failing_upload(folder_id, name, data, mime):
        if name.endswith(".pdf"):
            raise RuntimeError("Sensitive upload provider response")
        return original_upload(folder_id, name, data, mime)
    box.upload = failing_upload
    poll(box, store, clock, workflow)
    clock.now += 20
    with pytest.raises(NativeWorkflowError):
        poll(box, store, clock, workflow)
    assert store.value["report_ids"] == {"docx": "801"}
    assert poll(box, store, clock, workflow)["status"] == "failed"
    assert len(box.uploads) == 1 and workflow.calls == 1


def test_invalid_folder_config_never_reads_box_or_starts_a_workflow():
    box, store, clock, workflow = Box(), Store(), Clock(), Workflow()
    invalid = {**CONFIG, "reports_folder_id": "101"}
    with pytest.raises(NativeWorkflowError, match="configuration"):
        poll(box, store, clock, workflow, invalid)
    assert not box.listed and workflow.calls == 0


def test_review_export_shows_only_current_evaluation_and_evidence_mapping():
    import hashlib
    report = '# Annual review\n\nI checked the pilot measurement. [E1]'
    initial = '# Initial annual review\n\nI checked a measurement. [E1]'
    def evaluation(body, score, summary):
        return {'status': 'completed', 'report_hash': hashlib.sha256(body.encode()).hexdigest(), 'evaluated_at': '2026-10-17T08:00:00+00:00',
                'overall_score': score, 'summary': summary, 'strengths': ['Source-backed measurement.'], 'improvements': ['Explain cohort limits.'],
                'dimensions': [{'id': 'grounding', 'name': 'Source grounding', 'weight': 30, 'score': 4,
                                'rationale': 'The pilot claim cites the original measurement.', 'strengths': ['Exact source.'], 'improvements': ['Retain attribution.'], 'evidence_ids': ['E1']}],
                'leadership_alignment': [{'principle': 'Dive Deep', 'explanation': 'The draft describes checking the measurement.', 'evidence_ids': ['E1']}],
                'issues': [{'severity': 'review', 'message': 'Confirm cohort scope.', 'evidence_ids': ['E1']}], 'missing_information': ['What was the cohort?']}
    run = {'employee_name': 'Alex Chen', 'year': 2026, 'report_markdown': report, 'initial_draft': initial, 'reviewer_confirmed': False,
           'sources': [{'id': '201', 'name': 'Pilot.docx', 'url': 'https://app.box.com/file/201'}],
           'evidence': [{'id': 'E1', 'source_id': '201', 'source_name': 'Pilot.docx', 'claim': 'Checked the measurement.', 'quote': 'The pilot measurement was checked.'}],
           'evaluations': {'rubric_version': 'annual-review-v1', 'score_label': 'Draft quality', 'initial': evaluation(initial, 76, 'Initial draft finding.'),
                           'final': evaluation(report, 84, 'Final draft finding.'), 'overall_delta': 8}}
    notes = _review_document(run)
    markdown = notes['report_markdown']
    assert 'Current draft quality score: 84/100.' in markdown
    assert '76/100' not in markdown and 'change +8' not in markdown
    assert 'Initial draft finding.' not in markdown and 'Final draft finding.' in markdown
    assert 'Source grounding' in markdown and 'The pilot claim cites the original measurement.' in markdown
    assert 'Score: 4/5' in markdown
    assert 'Dive Deep: The draft describes checking the measurement. [E1]' in markdown
    assert 'Strength: Source-backed measurement.' in markdown and 'Improvement: Explain cohort limits.' in markdown
    assert 'Confirm cohort scope.' in markdown
    assert run['evaluations']['initial']['overall_score'] == 76
    assert 'SHA256' not in markdown
    assert run['evaluations']['initial']['report_hash'] not in markdown
    assert run['evaluations']['final']['report_hash'] not in markdown
    assert run['report_markdown'] == report
    from app.exports import make_export
    data, _, _ = make_export(notes, 'docx')
    exported = '\n'.join(p.text for p in Document(BytesIO(data)).paragraphs)
    assert 'Initial draft evaluation' not in exported
    assert '76/100' not in exported
    assert 'Current draft quality score: 84/100.' in exported
    run['report_markdown'] += '\n\nEmployee edits.'
    stale = _review_document(run)['report_markdown']
    assert 'No current score is available' in stale
    assert '84/100' not in stale and '76/100' not in stale


def test_old_review_export_does_not_invent_evaluations_or_scores():
    markdown = _review_document({'report_markdown': 'Historical draft.'})['report_markdown']
    assert 'No current evaluation was recorded' in markdown
    assert 'Current draft quality score:' not in markdown
