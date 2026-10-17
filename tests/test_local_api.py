"""Exercise the local workbench boundary without Box calls or model charges."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient
import pytest

from app.config import Settings
from app.main import create_app

CONFIG = {
    'root_folder_id': '100',
    'category_folders': {'Projects': '101', 'Peer Feedback': '102', 'Annual Plans': '103'},
    'reports_folder_id': '104', 'employee_name': 'Alex Chen', 'year': 2026, 'language': 'en',
}
PRIVATE = 'FULL SOURCE BODY MUST NOT REACH BROWSER OR PERSISTENCE'


class ImmediateExecutor:
    def submit(self, function, *args):
        function(*args)


class DeferredExecutor:
    def submit(self, function, *args):
        self.task = (function, args)


class Box:
    mode = 'client_credentials'

    def __init__(self):
        self.entries = {'101': [{'id': '201', 'type': 'file', 'name': 'Project.docx', 'size': 100, 'file_version': {'id': '1'}}], '102': [], '103': [], '104': []}
        self.uploads = []
        self.versions = []
        self.user_id = '42'
        self.fail_upload = False

    def identity(self):
        return {'id': self.user_id, 'name': 'Demo employee', 'type': 'user'}

    def list_folder(self, identifier):
        return {'entries': deepcopy(self.entries.get(identifier, [])), 'truncated': False}

    def upload(self, folder_id, name, data, mime):
        import hashlib
        from app.box_client import BoxError
        if self.fail_upload:
            raise BoxError('Box upload failed (HTTP 403). Check destination access.')
        identifier = str(800 + len(self.uploads))
        item = {'id': identifier, 'name': name, 'type': 'file', 'url': 'https://app.box.com/file/' + identifier,
                'size': len(data), 'extension': Path(name).suffix.lstrip('.'), 'file_version': {'id': '1'}, 'sha1': hashlib.sha1(data).hexdigest()}
        self.entries[folder_id].append(item)
        self.uploads.append({'folder_id': folder_id, 'data': data, **item})
        return item

    def file_info(self, identifier):
        from app.box_client import BoxError
        for folder_id, entries in self.entries.items():
            for item in entries:
                if item['id'] == identifier:
                    return {**deepcopy(item), 'type': 'file', 'parent': {'id': folder_id}, 'owned_by': {'id': self.user_id},
                            'etag': str(len(self.versions)), 'item_status': 'active'}
        raise BoxError('Not found.', status_code=404)

    def upload_version(self, identifier, name, data, mime, etag=None):
        from app.box_client import BoxError
        if self.fail_upload:
            raise BoxError('Box version upload failed.')
        assert etag is not None
        self.versions.append({'id': identifier, 'name': name, 'data': data})
        for entries in self.entries.values():
            for item in entries:
                if item['id'] == identifier:
                    item.update(name=name, size=len(data), file_version={'id': str(1 + len(self.versions))})
                    return {**item, 'url': 'https://app.box.com/file/' + identifier}
        raise BoxError('Not found.', status_code=404)


class Workflow:
    def __init__(self, box):
        self.box = box
        self.requests = []
        self.mutate = False

    def __call__(self, adapter, settings, request, update):
        self.requests.append(request)
        sources = [{**item, 'text': PRIVATE} for item in adapter.discover(request['folder_id'])['files']]
        update(stage='extract', progress=40, message='Extracting verified evidence.', sources=sources)
        if self.mutate:
            self.box.entries['101'][0]['file_version']['id'] = '2'
        return {'sources': sources, 'evidence': [{'id': 'E1', 'source_id': '201', 'source_name': 'Project.docx', 'quote': 'Measured pilot result.', 'claim': 'Team result', 'source_url': 'https://app.box.com/file/201'}],
                'report_markdown': '# Annual review\n\nI contributed to the project. [E1]', 'initial_draft': 'Initial draft. [E1]',
                'review': {'issues': [], 'revision_performed': True, 'box_ai': {'status': 'completed', 'file_ids': ['201'], 'citations': [{'id': '201', 'name': 'Project.docx', 'content': 'provider content'}]}},
                'coverage': {'selected': len(sources), 'skipped': [{'id': '299', 'name': 'Ignored.txt', 'reason': 'Unsupported format.'}]},
                'missing_information': ['What did you learn?'], 'usage': {'input_tokens': 100, 'output_tokens': 40}}


class Improvement:
    def __init__(self):
        self.calls, self.fail, self.same_body, self.bad_hash, self.mutate = [], False, False, False, False
        self.score = 78

    def __call__(self, adapter, settings, run, answers, target_score, update):
        import hashlib
        from app.workflow import WorkflowError
        self.calls.append((deepcopy(run), answers, target_score))
        update(stage='improve', progress=50, message='Applying one bounded improvement.')
        if self.fail:
            raise WorkflowError('Safe model validation failure.', 'improvement_validation', 'Only a private validation reason.')
        source = run['improvement_context']['new_source']
        assert answers in source['text'] and source['document_category'] == 'Annual Plans'
        result = deepcopy(run)
        result.pop('improvement_context', None)
        evidence_id = 'E' + str(len(result['evidence']) + 1)
        result['sources'].append({key: value for key, value in source.items() if key != 'text'})
        result['evidence'].append({'id': evidence_id, 'source_id': source['id'], 'source_name': source['name'], 'source_url': source['url'], 'quote': answers, 'claim': 'Self-reported learning.'})
        if not self.same_body:
            result['report_markdown'] += '\n\nEmployee-supplied reflection: ' + answers + ' [' + evidence_id + ']'
        final = {'status': 'completed', 'report_hash': hashlib.sha256(result['report_markdown'].encode()).hexdigest(), 'overall_score': self.score,
                 'summary': 'Current report evaluated.', 'dimensions': [], 'leadership_alignment': [], 'issues': [], 'strengths': [], 'improvements': [], 'missing_information': []}
        if self.bad_hash:
            final['report_hash'] = 'bad'
        result['evaluations'] = {'rubric_version': 'test', 'score_label': 'Draft quality', 'initial': deepcopy(run.get('evaluations', {}).get('initial', final)), 'final': final, 'overall_delta': 0}
        result['missing_information'] = ['What is the next concrete development step?']
        result['improvement'] = {'changed_sections': ['Overview']}
        if self.mutate:
            adapter.box.entries['101'][0]['file_version']['id'] = 'changed'
        return result


def fake_export(run, format):
    return b'EXPORT', 'annual-review.' + format, 'application/octet-stream'


@pytest.fixture
def client(tmp_path):
    box = Box()
    workflow = Workflow(box)
    app = create_app(Settings(box_user_id='42', data_dir=tmp_path), box, CONFIG, tmp_path / 'runs', workflow, ImmediateExecutor(), Improvement())
    with patch('app.main.make_export', side_effect=fake_export), TestClient(app, base_url='http://127.0.0.1') as client:
        yield client, box, workflow, tmp_path



def report_version(api, run_id):
    """Existing successful-flow tests submit the version they just loaded."""
    return {'expected_report_hash': api.get('/api/runs/' + run_id).json()['report_hash']}


def start(client):
    response = client.post('/api/runs', json={'notes': 'Employee-supplied context only.'})
    assert response.status_code == 202, response.text
    assert response.json()['status'] == 'completed', response.text
    return response.json()


def test_scope_async_draft_persistence_and_current_saved_revision(client):
    api, box, workflow, tmp_path = client
    status = api.get('/api/status').json()
    assert status['box']['connected']
    assert status['workspace_folder_url'] == 'https://app.box.com/folder/100'
    assert status['reports_folder_url'] == 'https://app.box.com/folder/104'
    assert [category['name'] for category in status['categories']] == list(CONFIG['category_folders'])
    assert 'id' not in status['categories'][0]['files'][0]
    run = start(api)
    assert workflow.requests[0]['folder_id'] == '100'
    assert workflow.requests[0]['language'] == 'en'
    assert len(box.uploads) == 3
    assert all(item['folder_id'] == '104' for item in box.uploads)
    assert all('Draft' in item['name'] for item in box.uploads)
    assert not run['reviewer_confirmed']
    assert len(run['box_files']) == 3
    assert 'text' not in run['sources'][0]
    assert 'source_id' not in run['evidence'][0]
    assert run['review']['box_ai']['file_count'] == 1
    assert run['review']['box_ai']['citations'] == [{'name': 'Project.docx'}]
    assert 'id' not in run['coverage']['skipped'][0]
    persisted = tmp_path / 'runs' / (run['id'] + '.json')
    assert persisted.stat().st_mode & 0o777 == 0o600
    assert PRIVATE not in persisted.read_text()
    assert api.get(f"/api/runs/{run['id']}/exports/pdf").content == b'EXPORT'
    changed = api.put(f"/api/runs/{run['id']}/report", json={**report_version(api, run['id']), 'report_markdown': 'Edited claim. [E1]'}).json()
    assert changed['box_files'] == []
    assert len(changed['saved_file_history']) == 3
    assert api.post(f"/api/runs/{run['id']}/save", json={'reviewer_confirmed': True}).status_code == 409
    confirmed = api.post(f"/api/runs/{run['id']}/confirm", json=report_version(api, run['id'])).json()
    assert confirmed['reviewer_confirmed']
    saved = api.post(f"/api/runs/{run['id']}/save", json={}).json()
    assert len(saved['files']) == 3
    assert all('Confirmed' in item['name'] for item in saved['files'])
    assert len(box.uploads) == 6
    api.post(f"/api/runs/{run['id']}/save", json={})
    assert len(box.uploads) == 6
    reset = api.put(f"/api/runs/{run['id']}/report", json={**report_version(api, run['id']), 'report_markdown': 'Another edit. [E1]'}).json()
    assert not reset['reviewer_confirmed']


def test_uploads_and_employee_addendum_are_fixed_scope(client):
    api, box, workflow, _ = client
    invalid = api.post('/api/uploads', data={'category': '99999'}, files={'files': ('Wrong.docx', b'fake', 'application/octet-stream')})
    assert invalid.status_code == 400
    invalid = api.post('/api/uploads', data={'category': 'Projects'}, files={'files': ('Wrong.exe', b'fake', 'application/octet-stream')})
    assert invalid.status_code == 400
    assert not box.uploads
    uploaded = api.post('/api/uploads', data={'category': 'Peer Feedback'}, files={'files': ('Feedback.docx', b'fake', 'application/octet-stream')})
    assert uploaded.status_code == 200
    assert box.uploads[0]['folder_id'] == '102'
    run = start(api)
    renewed = api.post(f"/api/runs/{run['id']}/answers", json={**report_version(api, run['id']), 'answers': 'I learned to validate launch criteria earlier.'})
    assert renewed.status_code == 202, renewed.text
    assert renewed.json()['id'] == run['id']
    addendum = next(item for item in box.uploads if item['folder_id'] == '103')
    document = Document(BytesIO(addendum['data']))
    text = '\n'.join(p.text for p in document.paragraphs)
    assert 'Employee-supplied' in text
    assert 'not independent measurements' in text
    assert 'Review year: 2026' in text
    assert 'I learned to validate launch criteria earlier.' in text
    assert len(workflow.requests) == 1
    assert len(api.app.state.workbench.improvement_workflow.calls) == 1
    assert not renewed.json()['reviewer_confirmed']


def test_unchanged_report_put_preserves_human_confirmation_and_edit_state(client):
    api, _, _, _ = client
    run = start(api)
    confirmed = api.post(f"/api/runs/{run['id']}/confirm", json=report_version(api, run['id'])).json()
    service = api.app.state.workbench
    before = service.get(run['id'])
    response = api.put(f"/api/runs/{run['id']}/report", json={**report_version(api, run['id']), 'report_markdown': confirmed['report_markdown']})
    assert response.status_code == 200 and response.json()['reviewer_confirmed'] is True
    assert response.json()['confirmed_at'] == confirmed['confirmed_at']
    assert service.get(run['id']) == before
    assert 'user_edited' not in response.json()


@pytest.mark.parametrize('action', ['edit', 'confirm', 'improve', 'answers'])
def test_report_mutations_require_the_version_actually_seen_by_the_client(client, action):
    import hashlib
    api, box, _, tmp_path = client
    run = start(api)
    stale_hash = run['report_hash']
    assert stale_hash == hashlib.sha256(run['report_markdown'].encode()).hexdigest()
    later_body = run['report_markdown'] + '\n\nA correction written in the other tab.'
    edited = api.put(f"/api/runs/{run['id']}/report", json={'report_markdown': later_body, 'expected_report_hash': stale_hash}).json()
    assert edited['report_hash'] != stale_hash
    service = api.app.state.workbench
    before = service.get(run['id'])
    uploads_before = len(box.uploads)
    if action == 'edit':
        response = api.put(f"/api/runs/{run['id']}/report", json={'report_markdown': 'Outdated overwrite.', 'expected_report_hash': stale_hash})
    else:
        payload = {'expected_report_hash': stale_hash}
        if action in {'improve', 'answers'}:
            payload['answers'] = 'New context intended for the earlier draft.'
        response = api.post(f"/api/runs/{run['id']}/{action}", json=payload)
    assert response.status_code == 409 and 'another tab' in response.json()['detail']
    assert service.get(run['id']) == before
    assert len(box.uploads) == uploads_before and not box.versions
    assert not list((tmp_path / 'improvement-checkpoints').glob('*.json'))
    assert service.improvement_workflow.calls == []
    assert api.get('/api/runs').json()['runs'][0]['report_hash'] == edited['report_hash']
    confirmed = api.post(f"/api/runs/{run['id']}/confirm", json={'expected_report_hash': edited['report_hash']}).json()
    assert confirmed['reviewer_confirmed'] is True and confirmed['report_markdown'] == later_body


@pytest.mark.parametrize('action', ['edit', 'confirm', 'improve', 'answers'])
def test_report_version_is_required_not_silently_assumed(client, action):
    api, box, _, _ = client
    run = start(api)
    before = api.app.state.workbench.get(run['id'])
    if action == 'edit':
        response = api.put(f"/api/runs/{run['id']}/report", json={'report_markdown': 'No version supplied.'})
    else:
        response = api.post(f"/api/runs/{run['id']}/{action}", json={'answers': 'No version supplied.'} if action in {'improve', 'answers'} else {})
    assert response.status_code == 422
    assert api.app.state.workbench.get(run['id']) == before
    assert len(box.uploads) == 3 and not box.versions


def test_cross_site_scope_validation_and_identity_guards(client):
    api, box, _, _ = client
    assert api.post('/api/runs', json={}, headers={'Origin': 'https://evil.example'}).status_code == 403
    assert api.post('/api/runs', json={}, headers={'Sec-Fetch-Site': 'cross-site'}).status_code == 403
    assert api.get('/api/status', headers={'Host': 'evil.example'}).status_code == 400
    assert api.post('/api/runs', json={'folder_id': '999'}).status_code == 422
    box.user_id = '999'
    mismatch = api.post('/api/runs', json={})
    assert mismatch.status_code == 400
    assert 'different user' in mismatch.json()['detail']
    assert not box.uploads


@pytest.fixture
def hosted_client(tmp_path):
    settings = Settings(box_user_id='42', data_dir=tmp_path, demo_username='judge', demo_password='offline-demo-password',
                        allowed_hosts=('demo.example.test',), public_origin='https://demo.example.test', origin_token='offline-origin-token')
    box = Box()
    app = create_app(settings, box, CONFIG, tmp_path / 'runs', Workflow(box), ImmediateExecutor(), Improvement())
    with patch('app.main.make_export', side_effect=fake_export), TestClient(app, base_url='https://demo.example.test') as api:
        yield api, box, settings


def test_hosted_demo_authentication_protects_every_route_except_safe_health(hosted_client):
    import base64
    api, box, settings = hosted_client
    assert api.get('/healthz', headers={'Host': 'internal-health-probe'}).json() == {'status': 'ok'}
    assert api.head('/healthz', headers={'Host': 'internal-health-probe'}).status_code == 200
    assert not box.uploads
    for path in ('/', '/api/status', '/api/runs', '/docs', '/openapi.json', '/assets/example.js'):
        assert api.get(path).status_code == 403
        response = api.get(path, headers={'X-Arsia-Origin': settings.origin_token})
        assert response.status_code == 401 and response.headers['www-authenticate'].startswith('Basic ')
        assert settings.demo_password not in response.text and settings.origin_token not in response.text
    headers = {'X-Arsia-Origin': settings.origin_token}
    for supplied in ('Bearer fake', 'Basic !invalid!', 'Basic ' + base64.b64encode(b'judge:incorrect-password').decode()):
        assert api.get('/api/status', headers={**headers, 'Authorization': supplied}).status_code == 401
    assert api.get('/api/status', headers=headers, auth=(settings.demo_username, settings.demo_password)).status_code == 200
    assert api.get('/api/status', headers={**headers, 'Host': 'untrusted.example.test'}, auth=(settings.demo_username, settings.demo_password)).status_code == 400
    assert settings.demo_password not in repr(settings) and settings.origin_token not in repr(settings)


def test_hosted_demo_mutations_still_reject_cross_site_requests(hosted_client):
    api, box, settings = hosted_client
    headers = {'X-Arsia-Origin': settings.origin_token, 'Origin': settings.public_origin}
    auth = (settings.demo_username, settings.demo_password)
    assert api.post('/api/runs', headers={**headers, 'Origin': 'https://evil.example.test'}, auth=auth, json={}).status_code == 403
    assert api.post('/api/runs', headers={**headers, 'Sec-Fetch-Site': 'cross-site'}, auth=auth, json={}).status_code == 403
    assert not box.uploads
    response = api.post('/api/runs', headers=headers, auth=auth, json={})
    assert response.status_code == 202 and response.json()['status'] == 'completed'
    assert len(box.uploads) == 3
    run = response.json()
    assert api.get(f"/api/runs/{run['id']}/exports/docx", headers={'X-Arsia-Origin': settings.origin_token}).status_code == 401
    assert api.get(f"/api/runs/{run['id']}/exports/docx", headers=headers, auth=auth).content == b'EXPORT'


@pytest.mark.parametrize('configuration', [
    {'allowed_hosts': ('*',)},
    {'allowed_hosts': ('demo.example.test',)},
    {'demo_username': 'judge'},
    {'demo_password': 'offline-demo-password'},
    {'demo_username': 'judge', 'demo_password': 'too-short'},
    {'allowed_hosts': ('demo.example.test',), 'public_origin': 'http://demo.example.test'},
    {'allowed_hosts': ('demo.example.test',), 'public_origin': 'https://other.example.test'},
    {'allowed_hosts': ('demo.example.test',), 'public_origin': 'https://demo.example.test/path'},
    {'allowed_hosts': ('demo.example.test',), 'public_origin': 'https://demo.example.test', 'origin_token': ''},
])
def test_hosted_demo_configuration_fails_closed(tmp_path, configuration):
    from dataclasses import replace
    defaults = Settings(data_dir=tmp_path, demo_username='judge', demo_password='offline-demo-password', origin_token='offline-origin-token')
    if configuration.keys() & {'demo_username', 'demo_password'}:
        defaults = Settings(data_dir=tmp_path)
    with pytest.raises(ValueError):
        create_app(replace(defaults, **configuration), Box(), CONFIG, tmp_path / 'runs', executor=ImmediateExecutor())


def test_workspace_configuration_uses_the_persistent_data_directory(tmp_path):
    import json
    from app.main import Workbench
    (tmp_path / 'native-config.json').write_text(json.dumps(CONFIG))
    service = Workbench(Settings(data_dir=tmp_path), Box(), executor=ImmediateExecutor())
    assert service.config() == CONFIG


def test_source_change_blocks_publishing_and_confirmation(client):
    api, box, workflow, _ = client
    workflow.mutate = True
    response = api.post('/api/runs', json={})
    assert response.json()['status'] == 'failed'
    assert 'Source documents changed' in response.json()['error']
    assert not box.uploads
    workflow.mutate = False
    run = start(api)
    assert api.post(f"/api/runs/{run['id']}/confirm", json=report_version(api, run['id'])).json()['reviewer_confirmed']
    box.entries['101'][0]['file_version']['id'] = '3'
    assert api.get(f"/api/runs/{run['id']}/exports/pdf").status_code == 409
    assert not api.get(f"/api/runs/{run['id']}").json()['reviewer_confirmed']
    assert api.post(f"/api/runs/{run['id']}/confirm", json=report_version(api, run['id'])).status_code == 409
    assert api.post(f"/api/runs/{run['id']}/save", json={}).status_code == 409


def test_partial_save_failure_preserves_local_draft_and_exposes_warning(client):
    api, box, _, _ = client
    box.fail_upload = True
    run = start(api)
    assert run['box_files'] == []
    assert len(run['save_warnings']) == 3
    assert api.get(f"/api/runs/{run['id']}/exports/docx").content == b'EXPORT'


def test_only_one_generation_and_uploads_blocked_while_busy(tmp_path):
    box = Box()
    workflow = Workflow(box)
    executor = DeferredExecutor()
    app = create_app(Settings(box_user_id='42', data_dir=tmp_path), box, CONFIG, tmp_path / 'runs', workflow, executor)
    with TestClient(app, base_url='http://127.0.0.1') as api:
        first = api.post('/api/runs', json={})
        assert first.json()['status'] == 'queued'
        assert api.post('/api/runs', json={}).status_code == 409
        assert api.post('/api/uploads', data={'category': 'Projects'}, files={'files': ('x.docx', b'fake')}).status_code == 409
        assert not workflow.requests
        assert not box.uploads


def test_actual_evaluations_are_whitelisted_copied_and_stale_after_edit(client):
    import hashlib
    from app.main import _public, _save_key
    api, _, _, _ = client
    run = start(api)
    final_hash = hashlib.sha256(run['report_markdown'].encode()).hexdigest()
    dimension = {'id': 'grounding', 'name': 'Source grounding', 'weight': 30, 'score': 4,
                 'rationale': 'Metric is cited.', 'strengths': ['Useful citation.'], 'improvements': ['Clarify the pilot.'], 'evidence_ids': ['E1'],
                 'private_source_text': 'EVALUATION_SECRET'}
    evaluation = {'status': 'completed', 'report_hash': final_hash, 'evaluated_at': '2026-10-17T08:00:00+00:00',
                  'overall_score': 80, 'summary': 'Grounded draft.', 'strengths': ['Specific measurement.'], 'improvements': ['Improve clarity.'],
                  'dimensions': [dimension], 'leadership_alignment': [{'principle': 'Dive Deep', 'explanation': 'Checked the metric.', 'evidence_ids': ['E1'], 'provider_secret': 'EVALUATION_SECRET'}],
                  'issues': [], 'missing_information': [], 'provider_request': {'secret': 'EVALUATION_SECRET'}}
    evaluations = {'rubric_version': 'annual-review-v1', 'score_label': 'Draft quality', 'initial': deepcopy(evaluation), 'final': deepcopy(evaluation),
                   'overall_delta': 0, 'credentials': 'EVALUATION_SECRET'}
    service = api.app.state.workbench
    service.update(run['id'], evaluations=evaluations)
    private = service.get(run['id'])
    public = _public(private)
    assert 'EVALUATION_SECRET' not in str(public)
    assert public['evaluations']['final']['status'] == 'completed'
    public['evaluations']['final']['strengths'].append('Only in copied response')
    assert 'Only in copied response' not in private['evaluations']['final']['strengths']
    new_score = deepcopy(private)
    new_score['evaluations']['final']['overall_score'] = 84
    assert _save_key(new_score) != _save_key(private)
    edited = api.put(f"/api/runs/{run['id']}/report", json={**report_version(api, run['id']), 'report_markdown': 'Employee revised the explanation. [E1]'}).json()
    assert edited['evaluations']['final']['status'] == 'stale'
    assert edited['evaluations']['overall_delta'] is None
    assert not edited['reviewer_confirmed']
    assert service.get(run['id'])['evaluations']['final']['status'] == 'completed'


def test_legacy_run_does_not_infer_evaluation_scores(client):
    api, _, _, _ = client
    run = start(api)
    assert 'evaluations' not in run


def test_failed_workflow_validation_diagnostics_are_private_and_bounded(client):
    import json
    from app.workflow import WorkflowError
    api, box, _, tmp_path = client
    error = WorkflowError('The report failed source validation. Retry with fewer documents.', 'model_output')
    error.validation_reason = 'INTERNAL_VALIDATION_DETAIL: required section was missing. ' * 30
    def fail(*_args):
        raise error
    api.app.state.workbench.workflow = fail
    response = api.post('/api/runs', json={})
    assert response.status_code == 202
    public = response.json()
    assert public['status'] == 'failed'
    assert 'internal_diagnostics' not in public
    assert 'INTERNAL_VALIDATION_DETAIL' not in response.text
    assert 'INTERNAL_VALIDATION_DETAIL' not in api.get(f"/api/runs/{public['id']}").text
    assert 'INTERNAL_VALIDATION_DETAIL' not in api.get('/api/runs').text
    private = json.loads((tmp_path / 'runs' / (public['id'] + '.json')).read_text())
    assert private['internal_diagnostics'] == {'code': 'model_output', 'validation_reason': error.validation_reason[:500]}
    assert len(private['internal_diagnostics']['validation_reason']) == 500
    assert not box.uploads


@pytest.mark.parametrize('status', ['queued', 'running'])
@pytest.mark.parametrize('stage', ['discover', 'read', 'extract'])
def test_public_hides_uncurated_extraction_candidates_but_private_state_keeps_them(status, stage):
    from app.main import _public
    evidence = [{'id': f'E{index}', 'source_id': '201', 'quote': f'Candidate source quotation {index}.'} for index in range(1, 51)]
    private = {'status': status, 'stage': stage, 'evidence': deepcopy(evidence)}
    assert _public(private)['evidence'] == []
    assert private['evidence'] == evidence


def test_public_exposes_selected_draft_evidence_without_truncating_legacy_references():
    from app.main import _public
    selected = [{'id': f'E{index}', 'source_id': '201', 'quote': f'Selected source quotation {index}.'} for index in range(1, 21)]
    draft = _public({'status': 'running', 'stage': 'draft', 'evidence': selected})
    assert [item['id'] for item in draft['evidence']] == [item['id'] for item in selected]
    assert 'source_id' not in draft['evidence'][0]
    legacy = [{'id': f'E{index}', 'source_id': '201', 'quote': f'Historical source quotation {index}.'} for index in range(1, 51)]
    completed = _public({'status': 'completed', 'stage': 'completed', 'evidence': legacy, 'report_markdown': 'Historical reference. [E47]'})
    assert len(completed['evidence']) == 50
    assert any(item['id'] == 'E47' for item in completed['evidence'])


class ImportBox(Box):
    def __init__(self):
        super().__init__()
        self.copies, self.copy_failures = [], {}
        self.entries['0'] = [{'id': '500', 'type': 'folder', 'name': 'External projects'},
                             {'id': '100', 'type': 'folder', 'name': 'Current review'},
                             {'id': '104', 'type': 'folder', 'name': 'Reports'},
                             {'id': '599', 'type': 'file', 'name': 'Image.png', 'size': 20, 'file_version': {'id': '1'}, 'extension': 'png'}]
        self.entries['500'] = []
        self.infos, self.versions = {}, {}
        for identifier, name, version in [('501', 'Design.docx', '11'), ('502', 'Results.pdf', '12'), ('503', 'Followup.docx', '13')]:
            info = {'id': identifier, 'type': 'file', 'name': name, 'size': 100, 'extension': Path(name).suffix.lstrip('.'),
                    'file_version': {'id': version}, 'path_collection': {'entries': [{'id': '0'}, {'id': '500'}]},
                    'expiring_embed_link': {'access_token': 'NEVER_EXPOSE_IMPORT_TOKEN'}}
            self.infos[identifier] = info
            self.entries['500'].append(info)
            self.versions[(identifier, version)] = ('PINNED SOURCE ' + version).encode()
        self.infos['201'] = {**deepcopy(self.entries['101'][0]), 'extension': 'docx', 'path_collection': {'entries': [{'id': '0'}, {'id': '100'}, {'id': '101'}]}}

    def list_folder(self, identifier):
        result = super().list_folder(identifier)
        result.update(id=identifier, name='All Files' if identifier == '0' else 'External projects',
                      parent=None if identifier == '0' else {'id': '0', 'name': 'All Files', 'access_token': 'NEVER_EXPOSE_IMPORT_TOKEN'},
                      path_collection={'entries': [{'id': '0'}]})
        if identifier == '600':
            result['path_collection']['entries'].append({'id': '100'})
        result['access_token'] = 'NEVER_EXPOSE_IMPORT_TOKEN'
        return result

    def file_info(self, identifier):
        return deepcopy(self.infos[identifier])

    def copy_file(self, file_id, folder_id, name, version_id):
        self.copies.append({'file_id': file_id, 'folder_id': folder_id, 'name': name, 'version_id': version_id})
        if file_id in self.copy_failures:
            raise self.copy_failures[file_id]
        copied = self.upload(folder_id, name, self.versions[(file_id, version_id)], 'application/octet-stream')
        self.infos[copied['id']] = {**deepcopy(copied), 'path_collection': {'entries': [{'id': '0'}, {'id': '100'}, {'id': folder_id}]}}
        return {**copied, 'expiring_embed_link': {'access_token': 'NEVER_EXPOSE_IMPORT_TOKEN'}}


@pytest.fixture
def import_client(tmp_path):
    box = ImportBox()
    app = create_app(Settings(box_user_id='42', data_dir=tmp_path), box, CONFIG, tmp_path / 'runs', Workflow(box), ImmediateExecutor())
    with TestClient(app, base_url='http://127.0.0.1') as api:
        yield api, box, tmp_path


def test_box_browse_is_safe_and_snapshot_copy_stays_independent(import_client):
    import json
    api, box, tmp_path = import_client
    root = api.get('/api/box/folders/0')
    assert root.status_code == 200
    assert [row['id'] for row in root.json()['entries']] == ['500']
    folder = api.get('/api/box/folders/500')
    assert folder.status_code == 200
    assert folder.json()['parent'] == {'id': '0', 'name': 'All Files'}
    assert folder.json()['entries'][0]['version_id'] == '11'
    assert 'NEVER_EXPOSE_IMPORT_TOKEN' not in root.text + folder.text
    imported = api.post('/api/imports', json={'category': 'Projects', 'files': [{'id': '501', 'version_id': '11'}]})
    assert imported.status_code == 200, imported.text
    assert imported.json()['warnings'] == []
    assert imported.json()['files'][0]['origin_url'] == 'https://app.box.com/file/501'
    assert 'NEVER_EXPOSE_IMPORT_TOKEN' not in imported.text
    assert box.copies == [{'file_id': '501', 'folder_id': '101', 'name': 'Design.docx', 'version_id': '11'}]
    manifest_path = tmp_path / 'imports.json'
    manifest = json.loads(manifest_path.read_text())
    assert manifest_path.stat().st_mode & 0o777 == 0o600
    assert manifest['imports'][0]['source_version_id'] == '11'
    assert manifest['imports'][0]['category'] == 'Projects'
    assert manifest['imports'][0]['mode'] == 'snapshot'
    box.infos['501']['file_version']['id'] = '111'
    box.versions[('501', '111')] = b'UPDATED ORIGINAL'
    assert box.uploads[0]['data'] == b'PINNED SOURCE 11'
    status = api.get('/api/status')
    materials = status.json()['categories'][0]['files']
    imported_row = next(row for row in materials if row['name'] == 'Design.docx')
    assert imported_row['origin_url'] == 'https://app.box.com/file/501'
    assert 'NEVER_EXPOSE_IMPORT_TOKEN' not in status.text


def test_import_preflight_blocks_changed_version_type_review_sources_and_foreign_destinations(import_client):
    api, box, _ = import_client
    assert api.get('/api/box/folders/100').status_code == 400
    assert api.get('/api/box/folders/600').status_code == 400
    assert api.post('/api/imports', json={'category': 'Projects', 'files': [{'id': '201', 'version_id': '1'}]}).status_code == 400
    assert api.post('/api/imports', json={'category': 'Projects', 'files': [{'id': '501', 'version_id': '10'}]}).status_code == 409
    box.infos['502']['extension'] = 'exe'
    assert api.post('/api/imports', json={'category': 'Projects', 'files': [{'id': '502', 'version_id': '12'}]}).status_code == 400
    box.infos['502']['extension'] = 'pdf'
    box.infos['502']['type'] = 'folder'
    assert api.post('/api/imports', json={'category': 'Projects', 'files': [{'id': '502', 'version_id': '12'}]}).status_code == 400
    assert api.post('/api/imports', json={'category': 'Other', 'files': [{'id': '501', 'version_id': '11'}]}).status_code == 422
    assert api.post('/api/imports', json={'category': 'Projects', 'destination_folder_id': '999', 'files': [{'id': '501', 'version_id': '11'}]}).status_code == 422
    assert api.post('/api/imports', json={'category': 'Projects', 'files': [{'id': '501', 'version_id': '11'}, {'id': '503', 'version_id': 'OLD'}]}).status_code == 422
    assert not box.copies


def test_import_checks_every_version_before_copying_and_rejects_wrong_identity(import_client):
    api, box, _ = import_client
    response = api.post('/api/imports', json={'category': 'Peer Feedback', 'files': [{'id': '501', 'version_id': '11'}, {'id': '502', 'version_id': '1'}]})
    assert response.status_code == 409 and not box.copies
    box.user_id = '999'
    assert api.get('/api/box/folders/0').status_code == 400
    assert api.post('/api/imports', json={'category': 'Projects', 'files': [{'id': '501', 'version_id': '11'}]}).status_code == 400
    assert not box.copies


def test_import_partial_and_ambiguous_outcomes_are_explicit_and_not_retried(import_client):
    import json
    from app.box_client import BoxError
    api, box, tmp_path = import_client
    box.copy_failures['502'] = BoxError('A file with this name already exists. No file was overwritten.', status_code=409)
    partial = api.post('/api/imports', json={'category': 'Annual Plans', 'files': [{'id': '501', 'version_id': '11'}, {'id': '502', 'version_id': '12'}]})
    assert partial.status_code == 200
    assert len(partial.json()['files']) == 1 and len(partial.json()['warnings']) == 1
    assert len(json.loads((tmp_path / 'imports.json').read_text())['imports']) == 1
    assert all(call['folder_id'] == '103' for call in box.copies)
    box.copy_failures['502'] = BoxError('Snapshot outcome uncertain; the copy was not retried.', ambiguous=True)
    count = len(box.copies)
    uncertain = api.post('/api/imports', json={'category': 'Projects', 'files': [{'id': '502', 'version_id': '12'}, {'id': '503', 'version_id': '13'}]})
    assert uncertain.status_code == 200 and uncertain.json()['files'] == []
    assert len(box.copies) == count + 1
    assert any('Remaining selections were not copied' in warning for warning in uncertain.json()['warnings'])
    assert all(call['file_id'] != '503' for call in box.copies)


class DeletionBox(Box):
    def __init__(self):
        super().__init__()
        self.metadata, self.removed, self.deletions, self.delete_failures = {}, set(), [], {}

    def upload(self, folder_id, name, data, mime):
        item = super().upload(folder_id, name, data, mime)
        self.metadata[item['id']] = {**deepcopy(item), 'type': 'file', 'parent': {'id': folder_id},
                                     'owned_by': {'id': '42'}, 'etag': '0', 'item_status': 'active'}
        return item

    def file_info(self, identifier):
        from app.box_client import BoxError
        if identifier in self.removed:
            raise BoxError('The output was not found.', status_code=404)
        return deepcopy(self.metadata[identifier])

    def trashed_file_info(self, identifier):
        from app.box_client import BoxError
        if identifier not in self.removed:
            raise BoxError('Unavailable.', status_code=404)
        return {**deepcopy(self.metadata[identifier]), 'item_status': 'trashed'}

    def trash_file(self, identifier, etag):
        assert etag == self.metadata[identifier]['etag']
        self.deletions.append(identifier)
        if identifier in self.delete_failures:
            raise self.delete_failures[identifier]
        self.removed.add(identifier)
        self.entries['104'] = [item for item in self.entries['104'] if item['id'] != identifier]
        return {'already_absent': False}


@pytest.fixture
def deletion_client(tmp_path):
    box = DeletionBox()
    app = create_app(Settings(box_user_id='42', data_dir=tmp_path), box, CONFIG, tmp_path / 'runs', Workflow(box), ImmediateExecutor())
    with patch('app.main.make_export', side_effect=fake_export), TestClient(app, base_url='http://127.0.0.1') as api:
        yield api, box, tmp_path


def lose_first_output_response(api, box):
    """The server creates the file, then the client loses its response."""
    import hashlib
    from app.box_client import BoxError
    original = box.upload
    lost = []
    def upload(folder_id, name, data, mime):
        if folder_id == CONFIG['reports_folder_id'] and not lost:
            private = next(iter(api.app.state.workbench.runs.values()))
            intent = next(item for item in private['pending_uploads'].values() if item['name'] == name)
            assert intent['content_sha1'] == [hashlib.sha1(data).hexdigest()]
            created = original(folder_id, name, data, mime)
            lost.append(created['id'])
            raise BoxError('The upload response was lost; not retried.', ambiguous=True)
        return original(folder_id, name, data, mime)
    box.upload = upload
    return lost


@pytest.mark.parametrize('retry_action', ['save', 'delete'])
def test_uncertain_created_output_is_verified_tracked_and_deletable_without_another_post(deletion_client, retry_action):
    api, box, _ = deletion_client
    lost = lose_first_output_response(api, box)
    run = start(api)
    assert len(box.uploads) == 3 and len(run['box_files']) == 2 and run['save_warnings']
    assert 'pending_uploads' not in run
    if retry_action == 'save':
        saved = api.post(f"/api/runs/{run['id']}/save", json={}).json()
        assert len(saved['files']) == 3 and saved['run']['save_warnings'] == []
        assert api.app.state.workbench.get(run['id'])['pending_uploads'] == {}
        assert len(box.uploads) == 3, 'Recover the created file instead of attempting the same POST again.'
    deleted = api.delete(f"/api/runs/{run['id']}").json()
    assert deleted['deleted'] is True and deleted['removed_count'] == 3
    assert set(box.deletions) == {item['id'] for item in box.uploads} and lost[0] in box.deletions
    assert box.entries['101'][0]['id'] == '201'


def test_upload_recovery_keeps_prior_content_hashes_when_export_bytes_change(deletion_client):
    import hashlib
    from app.box_client import BoxError
    api, box, _ = deletion_client
    original_upload, attempts, renders = box.upload, [], []
    def export(document, format):
        renders.append(format)
        return f'Export attempt {len(renders)}'.encode(), 'report.' + format, 'application/octet-stream'
    def upload(folder_id, name, data, mime):
        if name.endswith('.docx') and '-Review-notes' not in name:
            attempts.append((folder_id, name, data, mime))
            raise BoxError('The upload outcome is uncertain.', ambiguous=True)
        return original_upload(folder_id, name, data, mime)
    box.upload = upload
    with patch('app.main.make_export', side_effect=export):
        run = start(api)
        api.post(f"/api/runs/{run['id']}/save", json={})
        assert len(attempts) == 2 and attempts[0][2] != attempts[1][2]
        pending = next(iter(api.app.state.workbench.get(run['id'])['pending_uploads'].values()))
        assert set(pending['content_sha1']) == {hashlib.sha1(attempt[2]).hexdigest() for attempt in attempts}
        original_upload(*attempts[0])  # The earlier request becomes visible later.
        saved = api.post(f"/api/runs/{run['id']}/save", json={}).json()
        assert len(saved['files']) == 3 and saved['run']['save_warnings'] == []
        assert len(attempts) == 2 and len(box.uploads) == 3


@pytest.mark.parametrize('changed', ['hash', 'owner', 'location', 'shared', 'source', 'truncated'])
def test_uncertain_output_recovery_never_adopts_an_unverified_same_name_file(deletion_client, changed):
    api, box, _ = deletion_client
    lost = lose_first_output_response(api, box)
    run = start(api)
    target = lost[0]
    service = api.app.state.workbench
    if changed == 'hash':
        box.metadata[target]['sha1'] = '0' * 40
    elif changed == 'owner':
        box.metadata[target]['owned_by']['id'] = '999'
    elif changed == 'location':
        box.metadata[target]['parent']['id'] = '999'
    elif changed == 'shared':
        service.runs['a' * 32] = {'id': 'a' * 32, 'status': 'completed', 'box_files': [{'id': target}]}
    elif changed == 'source':
        service.runs['a' * 32] = {'id': 'a' * 32, 'status': 'completed', 'sources': [{'id': target}]}
    elif changed == 'truncated':
        original_list = box.list_folder
        box.list_folder = lambda identifier: {**original_list(identifier), 'truncated': identifier == CONFIG['reports_folder_id']}
    saved = api.post(f"/api/runs/{run['id']}/save", json={}).json()
    assert len(saved['files']) == 2 and len(saved['run']['save_warnings']) == 1
    assert len(box.uploads) == 3 and service.get(run['id'])['pending_uploads']
    deleted = api.delete(f"/api/runs/{run['id']}").json()
    assert deleted['deleted'] is False and deleted['status'] == 'partial'
    assert deleted['remaining_count'] == 3 and not box.deletions
    assert service.get(run['id'])['pending_uploads']


def test_delete_removes_only_tracked_outputs_history_and_is_idempotent(deletion_client):
    api, box, tmp_path = deletion_client
    run = start(api)
    assert run['deletion_preview'] == {'tracked_output_count': 3, 'can_delete': True}
    untracked = {'id': '999', 'type': 'file', 'name': 'Unrelated.docx'}
    box.entries['104'].append(untracked)
    result = api.delete(f"/api/runs/{run['id']}")
    assert result.status_code == 200, result.text
    assert result.json()['deleted'] is True and result.json()['removed_count'] == 3
    assert result.json()['remaining_count'] == 0
    assert set(box.deletions) == {'800', '801', '802'}
    assert box.entries['101'][0]['id'] == '201'
    assert untracked in box.entries['104']
    assert api.get(f"/api/runs/{run['id']}").status_code == 404
    assert api.get('/api/runs').json()['runs'] == []
    assert not (tmp_path / 'runs' / (run['id'] + '.json')).exists()
    tombstone = tmp_path / 'deleted-runs' / (run['id'] + '.json')
    assert tombstone.stat().st_mode & 0o777 == 0o600
    again = api.delete(f"/api/runs/{run['id']}")
    assert again.json() == result.json()
    assert len(box.deletions) == 3


def test_partial_deletion_retains_history_and_retries_only_remaining_outputs(deletion_client):
    from app.box_client import BoxError
    api, box, tmp_path = deletion_client
    run = start(api)
    box.delete_failures['801'] = BoxError('PRIVATE_PROVIDER_RESPONSE', status_code=403)
    partial = api.delete(f"/api/runs/{run['id']}")
    assert partial.status_code == 200 and partial.json()['deleted'] is False
    assert partial.json()['status'] == 'partial' and partial.json()['remaining_count'] == 1
    assert partial.json()['removed_count'] == 2
    assert 'PRIVATE_PROVIDER_RESPONSE' not in partial.text
    assert (tmp_path / 'runs' / (run['id'] + '.json')).exists()
    assert api.get(f"/api/runs/{run['id']}").json()['deletion']['remaining_count'] == 1
    assert api.post(f"/api/runs/{run['id']}/save", json={}).status_code == 409
    assert api.get(f"/api/runs/{run['id']}/exports/pdf").status_code == 409
    box.delete_failures.clear()
    retried = api.delete(f"/api/runs/{run['id']}").json()
    assert retried['deleted'] and retried['removed_count'] == 3
    assert box.deletions.count('800') == 1 and box.deletions.count('802') == 1 and box.deletions.count('801') == 2


def test_delete_preserves_shared_outputs_and_blocks_source_or_wrong_location(deletion_client):
    api, box, _ = deletion_client
    run = start(api)
    service = api.app.state.workbench
    other_id = 'a' * 32
    other = {'id': other_id, 'status': 'failed', 'stage': 'failed', 'created_at': '2026', 'events': [], 'box_files': [deepcopy(service.get(run['id'])['box_files'][0])]}
    service.runs[other_id] = other
    service.persist(other)
    box.metadata['801']['parent'] = {'id': '101'}
    box.metadata['802']['owned_by'] = {'id': '999'}
    partial = api.delete(f"/api/runs/{run['id']}").json()
    assert not partial['deleted'] and partial['shared_preserved_count'] == 1 and partial['remaining_count'] == 2
    assert not box.deletions
    box.metadata['801']['parent'] = {'id': '104'}
    box.metadata['802']['owned_by'] = {'id': '42'}
    done = api.delete(f"/api/runs/{run['id']}").json()
    assert done['deleted'] and done['shared_preserved_count'] == 1
    assert '800' not in box.deletions
    assert api.get(f'/api/runs/{other_id}').status_code == 200
    assert api.delete(f'/api/runs/{other_id}').json()['deleted']
    assert box.deletions.count('800') == 1


def test_delete_never_deletes_a_tracked_source_and_rejects_wrong_identity(deletion_client):
    api, box, _ = deletion_client
    run = start(api)
    service = api.app.state.workbench
    private = service.get(run['id'])
    private['sources'].append({'id': '800', 'name': 'Protected source.docx'})
    service.update(run['id'], sources=private['sources'])
    response = api.delete(f"/api/runs/{run['id']}").json()
    assert not response['deleted'] and response['remaining_count'] == 1
    assert '800' not in box.deletions
    box.user_id = '999'
    count = len(box.deletions)
    assert api.delete(f"/api/runs/{run['id']}").status_code == 400
    assert len(box.deletions) == count


def test_delete_rejects_active_runs_and_cross_site_requests(tmp_path):
    box = DeletionBox()
    app = create_app(Settings(box_user_id='42', data_dir=tmp_path), box, CONFIG, tmp_path / 'runs', Workflow(box), DeferredExecutor())
    with TestClient(app, base_url='http://127.0.0.1') as api:
        run = api.post('/api/runs', json={}).json()
        assert not run['deletion_preview']['can_delete']
        assert api.delete(f"/api/runs/{run['id']}").status_code == 409
        assert api.delete(f"/api/runs/{run['id']}", headers={'Origin': 'https://evil.example'}).status_code == 403
        assert not box.deletions


def test_legacy_errors_are_normalized_only_in_public_copy():
    from app.main import _public
    original = 'The report could not pass source-citation validation after one retry. Retry with a smaller set of documents.'
    run = {'status': 'failed', 'error': original, 'events': [{'message': original}]}
    public = _public(run)
    assert public['error'] == 'The generated report did not pass validation. Your materials are unchanged. Try generating it again.'
    assert public['events'][0]['message'] == public['error']
    assert run['error'] == original and run['events'][0]['message'] == original


def test_missing_output_requires_verified_trash_before_history_removal(deletion_client):
    from app.box_client import BoxError
    api, box, _ = deletion_client
    run = start(api)
    original = box.file_info
    def unavailable(identifier):
        if identifier == '800':
            raise BoxError('Not found or inaccessible.', status_code=404)
        return original(identifier)
    box.file_info = unavailable
    partial = api.delete(f"/api/runs/{run['id']}").json()
    assert not partial['deleted'] and partial['remaining_count'] == 1
    assert partial['already_absent_count'] == 0
    assert any('removal could not be verified' in warning for warning in partial['warnings'])
    assert '800' not in box.deletions
    box.removed.add('800')
    verified = api.delete(f"/api/runs/{run['id']}").json()
    assert verified['deleted'] and verified['already_absent_count'] == 1
    assert '800' not in box.deletions


def test_empty_failed_history_can_be_deleted_offline_but_tracked_outputs_still_require_identity(deletion_client):
    from app.box_client import BoxError
    api, box, tmp_path = deletion_client
    service = api.app.state.workbench
    calls = []
    def offline_identity():
        calls.append('identity')
        raise BoxError('Box is offline.')
    def forbidden_file_call(*_args):
        pytest.fail('Local-only deletion must not call any Box file operation.')
    box.identity = offline_identity
    box.file_info = forbidden_file_call
    box.trashed_file_info = forbidden_file_call
    box.trash_file = forbidden_file_call
    empty_id = 'b' * 32
    empty = {'id': empty_id, 'status': 'failed', 'stage': 'failed', 'events': [], 'sources': [], 'box_files': []}
    service.runs[empty_id] = empty
    service.persist(empty)
    deleted = api.delete('/api/runs/' + empty_id)
    assert deleted.status_code == 200 and deleted.json()['deleted'] is True
    assert deleted.json()['removed_count'] == 0
    assert not (tmp_path / 'runs' / (empty_id + '.json')).exists()
    assert calls == []
    tracked_id = 'c' * 32
    tracked = {'id': tracked_id, 'status': 'failed', 'stage': 'failed', 'events': [], 'sources': [],
               'box_files': [{'id': '800', 'name': 'Tracked.docx', 'kind': 'report', 'format': 'docx'}]}
    service.runs[tracked_id] = tracked
    service.persist(tracked)
    blocked = api.delete('/api/runs/' + tracked_id)
    assert blocked.status_code == 400
    assert calls == ['identity']
    assert (tmp_path / 'runs' / (tracked_id + '.json')).exists()
    assert tracked_id in service.runs


def test_finish_prepare_and_real_attachment_formats_with_fake_box(tmp_path):
    from pypdf import PdfReader
    box = DeletionBox()
    app = create_app(Settings(box_user_id='42', data_dir=tmp_path), box, CONFIG, tmp_path / 'runs', Workflow(box), ImmediateExecutor())
    with TestClient(app, base_url='http://127.0.0.1') as api:
        run = start(api)
        assert not run['reviewer_confirmed']
        for format, mime in [('docx', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'), ('pdf', 'application/pdf')]:
            before_uploads = len(box.uploads)
            prepared = api.get(f"/api/runs/{run['id']}/exports/{format}?prepare=true")
            assert prepared.status_code == 200
            assert prepared.json()['download_url'] == f"/api/runs/{run['id']}/exports/{format}"
            assert prepared.json()['filename'].endswith('.' + format)
            assert prepared.json()['content_type'] == mime
            assert len(box.uploads) == before_uploads
            attachment = api.get(prepared.json()['download_url'])
            assert attachment.status_code == 200
            assert attachment.headers['content-type'] == mime
            assert attachment.headers['content-disposition'].startswith('attachment; filename="')
            if format == 'docx':
                assert len(Document(BytesIO(attachment.content)).paragraphs) > 0
            else:
                assert len(PdfReader(BytesIO(attachment.content)).pages) > 0
        assert not api.get(f"/api/runs/{run['id']}").json()['reviewer_confirmed']
        confirmed = api.post(f"/api/runs/{run['id']}/confirm", json=report_version(api, run['id'])).json()
        assert confirmed['reviewer_confirmed']
        saved = api.post(f"/api/runs/{run['id']}/save", json={})
        assert saved.status_code == 200 and len(saved.json()['files']) == 3
        assert all('Confirmed' in item['name'] for item in saved.json()['files'])
        box.entries['101'][0]['file_version']['id'] = 'changed'
        assert api.get(f"/api/runs/{run['id']}/exports/docx?prepare=true").status_code == 409


def test_failed_runs_reject_finish_routes_but_allow_local_history_deletion(deletion_client):
    api, box, _ = deletion_client
    failed_id = 'd' * 32
    service = api.app.state.workbench
    service.runs[failed_id] = {'id': failed_id, 'status': 'failed', 'stage': 'failed', 'events': [], 'box_files': [], 'sources': []}
    service.persist(service.runs[failed_id])
    assert api.get(f'/api/runs/{failed_id}/exports/docx?prepare=true').status_code == 409
    assert api.get(f'/api/runs/{failed_id}/exports/pdf').status_code == 409
    assert api.post(f'/api/runs/{failed_id}/confirm', json=report_version(api, failed_id)).status_code == 409
    assert api.post(f'/api/runs/{failed_id}/save', json={}).status_code == 409
    assert api.put(f'/api/runs/{failed_id}/report', json={**report_version(api, failed_id), 'report_markdown': 'Edited failed draft.'}).status_code == 409
    assert api.post(f'/api/runs/{failed_id}/answers', json={**report_version(api, failed_id), 'answers': 'Additional context.'}).status_code == 409
    assert not box.uploads and not box.deletions
    assert api.delete(f'/api/runs/{failed_id}').json()['deleted']


class MaterialBox(ImportBox):
    def __init__(self):
        super().__init__()
        self.removed, self.material_deletions, self.material_failures = set(), [], {}
        for identifier, info in self.infos.items():
            info.update(parent={'id': '101' if identifier == '201' else '500'}, owned_by={'id': '42'}, etag='0', item_status='active')
        self.entries['104'] = [{'id': '999', 'type': 'file', 'name': 'Retained-report.pdf', 'size': 100, 'extension': 'pdf', 'file_version': {'id': '1'}}]

    def copy_file(self, file_id, folder_id, name, version_id):
        copied = super().copy_file(file_id, folder_id, name, version_id)
        self.infos[copied['id']].update(parent={'id': folder_id}, owned_by={'id': '42'}, etag='0', item_status='active')
        return copied

    def file_info(self, identifier):
        from app.box_client import BoxError
        if identifier in self.removed:
            raise BoxError('Not found or inaccessible.', status_code=404)
        return super().file_info(identifier)

    def trashed_file_info(self, identifier):
        from app.box_client import BoxError
        if identifier not in self.removed:
            raise BoxError('Unavailable.', status_code=404)
        return {**deepcopy(self.infos[identifier]), 'item_status': 'trashed'}

    def trash_file(self, identifier, etag):
        assert etag == self.infos[identifier]['etag']
        self.material_deletions.append(identifier)
        if identifier in self.material_failures:
            raise self.material_failures[identifier]
        self.removed.add(identifier)
        for parent, entries in self.entries.items():
            self.entries[parent] = [item for item in entries if item['id'] != identifier]
        return {'already_absent': False}


@pytest.fixture
def material_client(tmp_path):
    box = MaterialBox()
    app = create_app(Settings(box_user_id='42', data_dir=tmp_path), box, CONFIG, tmp_path / 'runs', Workflow(box), ImmediateExecutor())
    with TestClient(app, base_url='http://127.0.0.1') as api:
        yield api, box, tmp_path


def test_material_snapshot_delete_preserves_original_origin_reports_and_review_state(material_client):
    api, box, tmp_path = material_client
    imported = api.post('/api/imports', json={'category': 'Projects', 'files': [{'id': '501', 'version_id': '11'}]})
    assert imported.status_code == 200
    status = api.get('/api/status').json()
    row = next(item for item in status['categories'][0]['files'] if item['name'] == 'Design.docx')
    assert row['material_id'] == '800' and row['origin_url'] == 'https://app.box.com/file/501'
    assert 'material_id' not in status['reports'][0]
    original = deepcopy(box.infos['501'])
    service = api.app.state.workbench
    historical_id = 'e' * 32
    historical = {'id': historical_id, 'status': 'completed', 'stage': 'completed', 'reviewer_confirmed': True,
                  'events': [], 'sources': [{'id': '800', 'name': 'Design.docx'}], 'box_files': []}
    service.runs[historical_id] = historical
    service.persist(historical)
    deleted = api.delete('/api/materials/800')
    assert deleted.status_code == 200 and deleted.json() == {'deleted': True, 'status': 'deleted', 'warnings': []}
    assert box.material_deletions == ['800']
    assert box.infos['501'] == original and '501' not in box.removed
    assert box.entries['104'][0]['id'] == '999'
    assert service.get(historical_id)['reviewer_confirmed'] is True
    assert api.delete('/api/materials/800').json()['deleted']
    assert box.material_deletions == ['800']
    record = tmp_path / 'material-deletions' / '800.json'
    assert record.stat().st_mode & 0o777 == 0o600
    assert api.delete('/api/materials/501').status_code == 400


def test_material_deletion_supports_nested_membership_and_rejects_external_reports_or_folders(material_client):
    api, box, _ = material_client
    box.entries['101'] = [{'id': '150', 'type': 'folder', 'name': 'Nested'}]
    box.entries['150'] = [deepcopy(box.infos['201'])]
    box.infos['201'].update(parent={'id': '150'}, path_collection={'entries': [{'id': '0'}, {'id': '100'}, {'id': '101'}, {'id': '150'}]})
    assert api.get('/api/status').json()['categories'][0]['files'][0]['material_id'] == '201'
    for identifier in ['501', '999', '150', '100']:
        assert api.delete('/api/materials/' + identifier).status_code == 404
    assert not box.material_deletions
    assert api.delete('/api/materials/201').json()['deleted']
    assert box.material_deletions == ['201']


def test_material_owner_location_identity_and_active_generation_are_checked(material_client):
    api, box, _ = material_client
    box.infos['201']['owned_by'] = {'id': '999'}
    assert api.delete('/api/materials/201').status_code == 409
    box.infos['201']['owned_by'] = {'id': '42'}
    box.infos['201']['path_collection'] = {'entries': [{'id': '0'}, {'id': '500'}]}
    box.infos['201']['parent'] = {'id': '500'}
    assert api.delete('/api/materials/201').status_code == 409
    box.user_id = '999'
    assert api.delete('/api/materials/201').status_code == 400
    box.user_id = '42'
    api.app.state.workbench.runs['f' * 32] = {'id': 'f' * 32, 'status': 'running'}
    assert api.delete('/api/materials/201').status_code == 409
    assert not box.material_deletions


def test_material_unverified_404_is_retained_until_exact_trash_confirmation(material_client):
    from app.box_client import BoxError
    api, box, _ = material_client
    original = box.file_info
    def unavailable(identifier):
        if identifier == '201':
            raise BoxError('PRIVATE_PROVIDER_DETAIL', status_code=404)
        return original(identifier)
    box.file_info = unavailable
    uncertain = api.delete('/api/materials/201')
    assert uncertain.status_code == 200 and uncertain.json()['deleted'] is False
    assert uncertain.json()['status'] == 'unavailable'
    assert 'PRIVATE_PROVIDER_DETAIL' not in uncertain.text
    assert not box.material_deletions
    box.removed.add('201')
    verified = api.delete('/api/materials/201').json()
    assert verified['deleted']
    assert not box.material_deletions


def test_restart_persists_interrupted_stage_and_keeps_request_notes_private(tmp_path):
    import json
    from app.main import Workbench, _public
    run_id = '1' * 32
    service = Workbench(Settings(box_user_id='42', data_dir=tmp_path), Box(), CONFIG, tmp_path / 'runs', executor=DeferredExecutor())
    service.persist({'id': run_id, 'status': 'running', 'stage': 'initial_evaluation', 'events': [], 'request_notes': 'Private recovery context.'})
    service.load()
    recovered = json.loads((tmp_path / 'runs' / (run_id + '.json')).read_text())
    assert recovered['status'] == 'failed' and recovered['stage'] == 'failed'
    assert recovered['interrupted_stage'] == 'initial_evaluation'
    assert recovered['request_notes'] == 'Private recovery context.'
    assert 'request_notes' not in _public(recovered)
    assert 'interrupted_stage' not in _public(recovered)


def test_normal_lifespan_shutdown_waits_for_owned_generation(tmp_path, monkeypatch, caplog):
    import json
    from threading import Event, Thread
    box = Box()
    base_workflow = Workflow(box)
    entered, finish, run_created, allow_exit, shutting_down = (Event() for _ in range(5))
    outcome, failures = {}, []

    def controlled_workflow(*args):
        entered.set()
        assert finish.wait(5), 'Test did not release the controlled workflow.'
        return base_workflow(*args)

    app = create_app(Settings(box_user_id='42', data_dir=tmp_path), box, CONFIG, tmp_path / 'runs', controlled_workflow)
    executor = app.state.workbench.executor
    original_shutdown = executor.shutdown

    def observed_shutdown(*, wait, cancel_futures):
        assert wait is True and cancel_futures is False
        shutting_down.set()
        original_shutdown(wait=wait, cancel_futures=cancel_futures)

    monkeypatch.setattr(executor, 'shutdown', observed_shutdown)

    def serve():
        try:
            with TestClient(app, base_url='http://127.0.0.1') as api:
                response = api.post('/api/runs', json={'notes': 'Keep this context for private recovery.'})
                outcome.update(response.json())
                run_created.set()
                assert allow_exit.wait(5)
        except BaseException as error:
            failures.append(error)

    thread = Thread(target=serve)
    with patch('app.main.make_export', side_effect=fake_export):
        thread.start()
        try:
            assert entered.wait(5) and run_created.wait(5)
            assert 'request_notes' not in outcome
            allow_exit.set()
            assert shutting_down.wait(5)
            assert thread.is_alive(), 'Lifespan must wait until the owned worker completes.'
        finally:
            allow_exit.set()
            finish.set()
            thread.join(5)
    assert not thread.is_alive() and not failures
    persisted = json.loads((tmp_path / 'runs' / (outcome['id'] + '.json')).read_text())
    assert persisted['status'] == 'completed'
    assert persisted['request_notes'] == 'Keep this context for private recovery.'
    assert len(box.uploads) == 3
    assert 'Waiting for active generation to finish before shutting down.' in caplog.text


def test_manual_improvement_keeps_run_and_box_file_ids_with_honest_lower_scores(client):
    import hashlib
    import json
    api, box, workflow, tmp_path = client
    initial = start(api)
    original_draft = initial['initial_draft']
    edited_body = initial['report_markdown'] + '\n\nMy carefully edited sentence remains unchanged.'
    api.put(f"/api/runs/{initial['id']}/report", json={**report_version(api, initial['id']), 'report_markdown': edited_body})
    api.post(f"/api/runs/{initial['id']}/confirm", json=report_version(api, initial['id']))
    service = api.app.state.workbench
    prior = service.get(initial['id'])
    initial_eval = {'status': 'completed', 'report_hash': hashlib.sha256(edited_body.encode()).hexdigest(), 'overall_score': 80}
    service.update(initial['id'], evaluations={'rubric_version': 'test', 'score_label': 'Draft quality', 'initial': initial_eval, 'final': initial_eval, 'overall_delta': 0})
    original_output_ids = {item['id'] for item in prior['box_files']}
    answer = 'I learned to validate launch criteria earlier.'
    response = api.post(f"/api/runs/{initial['id']}/improve", json={**report_version(api, initial['id']), 'answers': answer, 'target_score': 95})
    assert response.status_code == 202
    improved = response.json()
    assert improved['id'] == initial['id'] and improved['status'] == 'completed'
    assert edited_body in improved['report_markdown']
    assert improved['initial_draft'] == original_draft and improved['user_edited'] is True
    assert improved['reviewer_confirmed'] is False
    assert improved['evaluations']['initial']['overall_score'] == 80
    assert improved['evaluations']['final']['overall_score'] == 78
    assert improved['evaluations']['overall_delta'] == -2
    assert improved['evaluations']['final']['report_hash'] == hashlib.sha256(improved['report_markdown'].encode()).hexdigest()
    assert improved['improvement']['target_reached'] is False
    assert improved['missing_information'] == ['What is the next concrete development step?']
    assert len(workflow.requests) == 1 and len(api.get('/api/runs').json()['runs']) == 1
    assert len(box.versions) == 3 and {item['id'] for item in box.versions} == original_output_ids
    private = service.get(initial['id'])
    assert len(private['revision_audit']) == 1 and private['improvement_answers'][0]['answers'] == answer
    assert original_output_ids == {item['id'] for item in private['box_files']}
    assert 'revision_audit' not in improved and 'improvement_answers' not in improved and 'request_notes' not in improved
    checkpoints = list((tmp_path / 'improvement-checkpoints').glob('*.json'))
    assert len(checkpoints) == 1 and checkpoints[0].stat().st_mode & 0o777 == 0o600
    assert json.loads(checkpoints[0].read_text())['run']['report_markdown'] == edited_body
    service.improvement_workflow.score = 75
    again = api.post(f"/api/runs/{initial['id']}/answers", json={**report_version(api, initial['id']), 'answers': 'My next step is an earlier review checklist.'}).json()
    assert again['id'] == initial['id'] and again['improvement']['round'] == 2
    assert again['evaluations']['initial']['overall_score'] == 80 and again['evaluations']['final']['overall_score'] == 75
    assert len(box.versions) == 6 and {item['id'] for item in box.versions} == original_output_ids
    assert len(box.uploads) == 5  # Three initial outputs and two answer documents.
    assert len(workflow.requests) == 1


def test_reverting_report_republishes_reused_box_files_and_current_save_is_idempotent(client):
    api, box, _, _ = client
    run = start(api)
    improved = api.post(f"/api/runs/{run['id']}/improve", json={**report_version(api, run['id']), 'answers': 'I learned to define the launch checklist earlier.'}).json()
    first_body = improved['report_markdown']
    original_ids = {item['id'] for item in box.versions}
    assert len(box.versions) == 3

    def content_export(document, format):
        return document['report_markdown'].encode(), 'annual-review.' + format, 'application/octet-stream'

    with patch('app.main.make_export', side_effect=content_export):
        second_body = first_body + '\n\nMy later human edit remains specific.'
        api.put(f"/api/runs/{run['id']}/report", json={**report_version(api, run['id']), 'report_markdown': second_body})
        saved = api.post(f"/api/runs/{run['id']}/save", json={}).json()
        assert len(saved['files']) == 3 and len(box.versions) == 6
        assert all(item['data'] == second_body.encode() for item in box.versions[-3:-1])

        reverted = api.put(f"/api/runs/{run['id']}/report", json={**report_version(api, run['id']), 'report_markdown': first_body}).json()
        assert reverted['box_files'] == [], 'Historical links now contain the later content and must not appear current.'
        saved = api.post(f"/api/runs/{run['id']}/save", json={}).json()
        assert len(saved['files']) == 3 and saved['run']['save_warnings'] == []
        assert len(box.versions) == 9 and {item['id'] for item in box.versions[-3:]} == original_ids
        assert all(item['data'] == first_body.encode() for item in box.versions[-3:-1])
        assert len(saved['run']['box_files']) == 3

        again = api.post(f"/api/runs/{run['id']}/save", json={}).json()
        assert len(again['files']) == 3 and len(box.versions) == 9


@pytest.mark.parametrize('failure_mode', ['model', 'evaluation_hash'])
def test_failed_improvement_restores_usable_report_and_confirmation(client, failure_mode):
    api, box, _, _ = client
    run = start(api)
    api.post(f"/api/runs/{run['id']}/confirm", json=report_version(api, run['id']))
    service = api.app.state.workbench
    before = service.get(run['id'])
    service.improvement_workflow.fail = failure_mode == 'model'
    service.improvement_workflow.bad_hash = failure_mode == 'evaluation_hash'
    response = api.post(f"/api/runs/{run['id']}/improve", json={**report_version(api, run['id']), 'answers': 'I learned to request earlier review.'})
    assert response.status_code == 202
    restored = response.json()
    assert restored['id'] == run['id'] and restored['status'] == 'completed'
    assert restored['report_markdown'] == before['report_markdown'] and restored['reviewer_confirmed'] is True
    assert restored['initial_draft'] == before['initial_draft']
    assert restored['improvement']['status'] == 'failed' and 'previous report has been restored' in restored['improvement_error']
    assert 'Only a private validation reason.' not in str(restored)
    assert len(box.uploads) == 4 and not box.versions
    assert len(api.get('/api/runs').json()['runs']) == 1
    assert api.get(f"/api/runs/{run['id']}/exports/docx").status_code == 200
    assert service.get(run['id'])['pending_addendum']['id'] == box.uploads[-1]['id']


def test_improvement_no_body_change_keeps_confirmation_and_existing_edit_flag(client):
    api, box, _, _ = client
    run = start(api)
    api.put(f"/api/runs/{run['id']}/report", json={**report_version(api, run['id']), 'report_markdown': run['report_markdown'] + '\n\nHuman wording.'})
    before = api.post(f"/api/runs/{run['id']}/confirm", json=report_version(api, run['id'])).json()
    api.app.state.workbench.improvement_workflow.same_body = True
    response = api.post(f"/api/runs/{run['id']}/improve", json={**report_version(api, run['id']), 'answers': 'This answer adds no new report wording.'}).json()
    assert response['report_markdown'] == before['report_markdown']
    assert response['reviewer_confirmed'] is True and response['confirmed_at'] == before['confirmed_at']
    assert response['user_edited'] is True and response['improvement']['report_changed'] is False
    assert len(box.versions) == 3


def test_improvement_source_changes_block_start_and_publish(client):
    api, box, _, _ = client
    run = start(api)
    uploads_before = len(box.uploads)
    box.entries['101'][0]['file_version']['id'] = '2'
    assert api.post(f"/api/runs/{run['id']}/improve", json={**report_version(api, run['id']), 'answers': 'A concrete learning.'}).status_code == 409
    assert len(box.uploads) == uploads_before
    box.entries['101'][0]['file_version']['id'] = '1'
    api.app.state.workbench.improvement_workflow.mutate = True
    restored = api.post(f"/api/runs/{run['id']}/improve", json={**report_version(api, run['id']), 'answers': 'A concrete learning.'}).json()
    assert restored['status'] == 'completed' and restored['report_markdown'] == run['report_markdown']
    assert restored['improvement']['status'] == 'failed' and not box.versions
    assert api.get(f"/api/runs/{run['id']}/exports/pdf").status_code == 409


def test_async_improvement_keeps_one_history_row_and_recovers_previous_checkpoint(client):
    import json
    api, box, _, tmp_path = client
    run = start(api)
    service = api.app.state.workbench
    deferred = DeferredExecutor()
    service.executor = deferred
    response = api.post(f"/api/runs/{run['id']}/improve", json={**report_version(api, run['id']), 'answers': 'Private pending answer.'})
    assert response.status_code == 202
    pending = response.json()
    assert pending['id'] == run['id'] and pending['status'] == 'running' and pending['stage'] == 'improve'
    assert pending['report_markdown'] == run['report_markdown']
    assert api.post(f"/api/runs/{run['id']}/improve", json={**report_version(api, run['id']), 'answers': 'Another round.'}).status_code == 409
    assert api.delete('/api/materials/201').status_code == 409
    assert api.delete('/api/runs/' + run['id']).status_code == 409
    assert len(box.uploads) == 3  # The deferred worker has not saved answers yet.
    service.load()
    restored = api.get('/api/runs/' + run['id']).json()
    assert restored['status'] == 'completed' and restored['report_markdown'] == run['report_markdown']
    assert restored['improvement']['status'] == 'failed'
    persisted = json.loads((tmp_path / 'runs' / (run['id'] + '.json')).read_text())
    assert persisted['interrupted_stage'] == 'improve' and 'Private pending answer.' in persisted['request_notes']
    assert len(api.get('/api/runs').json()['runs']) == 1


def test_improvement_partial_box_version_failure_retains_new_local_report_for_retry(client):
    api, box, _, _ = client
    run = start(api)
    original_upload_version = box.upload_version
    calls = []
    def fail_one(identifier, *args, **kwargs):
        from app.box_client import BoxError
        calls.append(identifier)
        if len(calls) == 2:
            raise BoxError('A safe version-upload failure.')
        return original_upload_version(identifier, *args, **kwargs)
    box.upload_version = fail_one
    result = api.post(f"/api/runs/{run['id']}/improve", json={**report_version(api, run['id']), 'answers': 'I learned to review the criteria earlier.'}).json()
    assert result['status'] == 'completed' and result['report_markdown'] != run['report_markdown']
    assert len(result['save_warnings']) == 1 and len(result['box_files']) == 2
    saved = api.post(f"/api/runs/{run['id']}/save", json={}).json()
    assert len(saved['files']) == 3 and saved['run']['save_warnings'] == []
    assert len(box.uploads) == 4 and len(box.versions) == 3


@pytest.mark.parametrize('change', ['owner', 'folder', 'shared'])
def test_improvement_never_overwrites_foreign_or_shared_outputs(client, change):
    api, box, _, _ = client
    run = start(api)
    service = api.app.state.workbench
    target = service.get(run['id'])['box_files'][0]['id']
    original_info = box.file_info
    def changed_info(identifier):
        info = original_info(identifier)
        if identifier == target:
            if change == 'owner':
                info['owned_by'] = {'id': '999'}
            elif change == 'folder':
                info['parent'] = {'id': '999'}
        return info
    box.file_info = changed_info
    if change == 'shared':
        service.runs['9' * 32] = {'id': '9' * 32, 'status': 'completed', 'box_files': [{'id': target}]}
    improved = api.post(f"/api/runs/{run['id']}/improve", json={**report_version(api, run['id']), 'answers': 'I learned to ask for earlier review.'}).json()
    assert improved['status'] == 'completed' and len(improved['save_warnings']) == 1
    assert target not in {item['id'] for item in box.versions}
    assert len(box.versions) == 2


def test_answer_upload_failure_does_not_claim_the_answer_was_saved_to_box(client):
    api, box, _, _ = client
    run = start(api)
    box.fail_upload = True
    result = api.post(f"/api/runs/{run['id']}/improve", json={**report_version(api, run['id']), 'answers': 'I learned to review launch criteria earlier.'}).json()
    assert result['status'] == 'completed' and result['report_markdown'] == run['report_markdown']
    assert 'private local history' in result['improvement_error']
    assert 'saved as self-reported material' not in result['improvement_error']
    assert len(box.uploads) == 3 and not box.versions


@pytest.mark.parametrize('verified', [True, False])
def test_lost_addendum_response_requires_verified_content_before_improvement(client, verified):
    from app.box_client import BoxError
    api, box, _, _ = client
    run = start(api)
    original_upload, answer_posts = box.upload, []
    def upload(folder_id, name, data, mime):
        created = original_upload(folder_id, name, data, mime)
        if folder_id == CONFIG['category_folders']['Annual Plans']:
            answer_posts.append(created['id'])
            if not verified:
                next(item for item in box.entries[folder_id] if item['id'] == created['id'])['sha1'] = '0' * 40
            raise BoxError('The answer-upload response was lost; not retried.', ambiguous=True)
        return created
    box.upload = upload
    result = api.post(f"/api/runs/{run['id']}/improve", json={**report_version(api, run['id']), 'answers': 'I learned to review the task definition earlier.'}).json()
    assert result['id'] == run['id'] and result['status'] == 'completed'
    assert len(answer_posts) == 1 and len(box.uploads) == 4
    private = api.app.state.workbench.get(run['id'])
    if verified:
        assert result['improvement']['status'] == 'completed' and result['report_markdown'] != run['report_markdown']
        assert len(box.versions) == 3 and private['addendum_upload_intent'] is None
    else:
        assert result['improvement']['status'] == 'failed' and result['report_markdown'] == run['report_markdown']
        assert 'Check Materials' in result['improvement_error'] and 'remove that unapplied answer copy' in result['improvement_error']
        assert private['addendum_upload_intent']['status'] == 'uncertain' and 'addendum_upload_intent' not in result
        assert api.app.state.workbench.improvement_workflow.calls == [] and not box.versions
        assert api.post(f"/api/runs/{run['id']}/improve", json={**report_version(api, run['id']), 'answers': 'Another attempt.'}).status_code == 409


def test_restart_after_first_box_version_keeps_committed_candidate_and_retries_missing_outputs(client):
    import hashlib
    api, box, _, _ = client
    run = start(api)
    service = api.app.state.workbench
    deferred = DeferredExecutor()
    service.executor = deferred
    response = api.post(f"/api/runs/{run['id']}/improve", json={**report_version(api, run['id']), 'answers': 'I learned to validate the task definition earlier.'})
    assert response.json()['status'] == 'running'
    original_upload_version = box.upload_version
    calls = []
    def terminate_second_version(*args, **kwargs):
        calls.append(args[0])
        if len(calls) == 2:
            raise SystemExit('Simulated abrupt process interruption, not a provider error.')
        return original_upload_version(*args, **kwargs)
    box.upload_version = terminate_second_version
    function, args = deferred.task
    with pytest.raises(SystemExit):
        function(*args)
    committed = service.get(run['id'])
    assert committed['status'] == 'running' and committed['stage'] == 'save_draft'
    assert committed['improvement']['content_committed'] is True
    assert committed['report_markdown'] != run['report_markdown'] and len(box.versions) == 1
    expected_hash = hashlib.sha256(committed['report_markdown'].encode()).hexdigest()
    service.load()
    recovered = api.get('/api/runs/' + run['id']).json()
    assert recovered['status'] == 'completed' and recovered['report_markdown'] == committed['report_markdown']
    assert recovered['evaluations']['final']['report_hash'] == expected_hash
    assert recovered['evaluations']['final']['status'] == 'completed'
    assert recovered['save_warnings'] and 'current report is retained' in recovered['improvement_error']
    assert len(recovered['box_files']) == 1
    box.upload_version = original_upload_version
    saved = api.post('/api/runs/' + run['id'] + '/save', json={}).json()
    assert len(saved['files']) == 3 and saved['run']['save_warnings'] == []
    assert len(box.versions) == 3 and len(box.uploads) == 4
