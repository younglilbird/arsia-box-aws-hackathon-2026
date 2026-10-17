"""Small bounded Box adapter. Tokens stay in server memory."""

from collections import deque
from io import BytesIO
import json
from pathlib import Path
import re
from threading import Lock
import time
from urllib.parse import urlsplit
from zipfile import ZipFile, BadZipFile

import httpx

from app.config import Settings

MAX_FILE_BYTES = 12 * 1024 * 1024
EXTENSIONS = {"txt", "md", "csv", "pdf", "docx", "pptx"}


class BoxError(RuntimeError):
    """Safe, user-facing integration error."""

    def __init__(self, message: str, *, status_code: int | None = None, ambiguous: bool = False):
        super().__init__(message)
        self.status_code, self.ambiguous = status_code, ambiguous


def box_id(value: str) -> str:
    if not re.fullmatch(r"[0-9]{1,30}", str(value)):
        raise BoxError("Box file and folder IDs must contain digits only.")
    return str(value)


def check_archive(data: bytes) -> None:
    try:
        with ZipFile(BytesIO(data)) as archive:
            if len(archive.infolist()) > 3000 or sum(i.file_size for i in archive.infolist()) > 40 * 1024 * 1024:
                raise BoxError("This Office document is too large when expanded. Use a smaller file.")
    except BadZipFile:
        raise BoxError("This Office document is not a valid readable archive.") from None


def extract_text(data: bytes, extension: str) -> str:
    if len(data) > MAX_FILE_BYTES:
        raise BoxError("The file exceeds 12 MB. Use a smaller file.")
    if extension in {"txt", "md", "csv"}:
        for encoding in ("utf-8-sig", "utf-16", "gb18030"):
            try:
                return data.decode(encoding).replace("\x00", "")
            except UnicodeError:
                continue
        raise BoxError("Cannot decode the text. Save it as UTF-8 and retry.")
    if extension == "pdf":
        from pypdf import PdfReader
        document = PdfReader(BytesIO(data))
        if document.is_encrypted:
            raise BoxError("Cannot read an encrypted PDF. Decrypt it first.")
        if len(document.pages) > 80:
            raise BoxError("This PDF exceeds 80 pages. Split it into smaller documents.")
        return "\n\n".join(page.extract_text() or "" for page in document.pages)
    if extension in {"docx", "pptx"}:
        check_archive(data)
    if extension == "docx":
        from docx import Document
        document = Document(BytesIO(data))
        chunks = [p.text for p in document.paragraphs]
        for table in document.tables:
            chunks.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(chunks)
    if extension == "pptx":
        from pptx import Presentation
        document = Presentation(BytesIO(data))
        chunks = []
        for index, slide in enumerate(document.slides, 1):
            chunks.append(f"Slide {index}")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    chunks.append(shape.text)
                if shape.has_table:
                    chunks.extend(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)
        return "\n".join(chunks)
    raise BoxError("This file format is not supported.")


class BoxClient:
    def __init__(self, settings: Settings):
        self.settings = settings
        self._token = ""
        self._expires = 0.0
        self._token_lock = Lock()
        self.http = httpx.Client(timeout=httpx.Timeout(65, connect=10), follow_redirects=False)

    @property
    def mode(self) -> str:
        return "client_credentials" if self.settings.box_client_id and self.settings.box_client_secret and self.settings.box_enterprise_id else "access_token"

    def token(self, force: bool = False) -> str:
        with self._token_lock:
            if self._token and self._expires > time.time() and not force:
                return self._token
            if self.mode == "client_credentials":
                try:
                    response = self.http.post("https://api.box.com/oauth2/token", data={
                        "grant_type": "client_credentials", "client_id": self.settings.box_client_id,
                        "client_secret": self.settings.box_client_secret,
                        "box_subject_type": "user" if self.settings.box_user_id else "enterprise",
                        "box_subject_id": self.settings.box_user_id or self.settings.box_enterprise_id,
                    })
                    if response.status_code != 200:
                        raise BoxError(f"Box authentication failed (HTTP {response.status_code}). Check the local .env file and application authorization.")
                    payload = response.json()
                    self._token = payload["access_token"]
                    self._expires = time.time() + max(0, int(payload.get("expires_in", 3600)) - 60)
                except (httpx.HTTPError, ValueError, KeyError):
                    raise BoxError("Cannot authenticate to Box. Check your network and local configuration.") from None
            elif self.settings.box_access_token:
                self._token = self.settings.box_access_token
                self._expires = time.time() + 60
            else:
                raise BoxError("Box is not configured. Add the application credentials to your local .env file.")
            return self._token

    def request(self, method: str, path: str, *, retry: bool = True, **kwargs) -> dict:
        base = "https://upload.box.com/api" if re.fullmatch(r"/2\.0/files(?:/[0-9]+)?/content", path) else "https://api.box.com"
        attempts = 3 if retry else 1
        extra_headers = kwargs.pop("headers", {})
        for attempt in range(attempts):
            try:
                response = self.http.request(method, base + path, headers={**extra_headers, "Authorization": "Bearer " + self.token()}, **kwargs)
            except httpx.HTTPError:
                raise BoxError("Cannot connect to Box. Check your network and retry.", ambiguous=not retry) from None
            if response.status_code == 401 and attempt == 0 and attempt + 1 < attempts and self.mode == "client_credentials":
                self.token(force=True)
                continue
            if response.status_code in {429, 502, 503, 504} and attempt + 1 < attempts:
                time.sleep(attempt + 1)
                continue
            if response.status_code >= 400:
                instructions = {401: "The token has expired. Sign in again or update the local configuration.", 403: "The application does not have access to this content.", 404: "The file or folder does not exist or is not accessible to this application.", 409: "A file with this name already exists. Choose a different name.", 429: "Too many requests. Please retry shortly."}
                raise BoxError(f"Box HTTP {response.status_code}: " + instructions.get(response.status_code, "The operation did not complete. Please retry shortly."), status_code=response.status_code, ambiguous=not retry and response.status_code >= 500)
            if response.status_code == 204:
                return {}
            try:
                return response.json()
            except ValueError:
                raise BoxError("Cannot parse the response from Box.", ambiguous=not retry) from None
        raise BoxError("Box did not complete the request after bounded retries.")

    def identity(self) -> dict:
        identity = self.request("GET", "/2.0/users/me", params={"fields": "id,name,type"})
        return {key: identity.get(key) for key in ("id", "name", "type")}

    def list_folder(self, folder_id: str) -> dict:
        folder_id = box_id(folder_id)
        info = self.request("GET", f"/2.0/folders/{folder_id}", params={"fields": "id,name,parent,path_collection"})
        entries, marker = [], None
        for _ in range(30):
            params = {"limit": 200, "usemarker": "true", "fields": "id,type,name,size,extension,file_version"}
            if marker:
                params["marker"] = marker
            page = self.request("GET", f"/2.0/folders/{folder_id}/items", params=params)
            entries.extend({**item, "url": f"https://app.box.com/{'folder' if item['type'] == 'folder' else 'file'}/{item['id']}"} for item in page.get("entries", []))
            marker = page.get("next_marker")
            if not marker:
                break
        return {"id": folder_id, "name": info.get("name", "Box"), "parent": info.get("parent"), "path_collection": info.get("path_collection"), "entries": entries, "truncated": bool(marker)}

    def file_info(self, file_id: str) -> dict:
        return self.request("GET", f"/2.0/files/{box_id(file_id)}", params={"fields": "id,type,name,extension,size,file_version,path_collection,parent,owned_by,etag,item_status"})

    def trashed_file_info(self, file_id: str) -> dict:
        return self.request("GET", f"/2.0/files/{box_id(file_id)}/trash", params={"fields": "id,type,name,item_status,owned_by"})

    def trash_file(self, file_id: str, etag: str) -> dict:
        """Normal Box deletion only; enterprise Trash policy controls recovery."""
        identifier = box_id(file_id)
        if not isinstance(etag, str) or not re.fullmatch(r"[A-Za-z0-9._:-]{1,120}", etag):
            raise BoxError("The output version could not be verified for deletion. Refresh and try again.")
        try:
            self.request("DELETE", f"/2.0/files/{identifier}", retry=False, headers={"If-Match": etag})
        except BoxError as error:
            if error.status_code == 404:
                return {"already_absent": False, "needs_verification": True}
            if error.status_code == 412:
                raise BoxError("The Box output changed before deletion. Nothing was deleted by this request; retry after checking the file.", status_code=412) from None
            if error.ambiguous:
                raise BoxError("Box did not confirm the deletion outcome. The draft history was retained so deletion can be retried safely.", ambiguous=True) from None
            raise
        return {"already_absent": False}

    def copy_file(self, file_id: str, folder_id: str, name: str, version_id: str) -> dict:
        """Copy exactly one pinned version; never retry an ambiguous mutation."""
        source_id, destination_id, version_id = box_id(file_id), box_id(folder_id), box_id(version_id)
        if not name or len(name) > 255 or any(c in name for c in ("/", "\\", "\r", "\n", "\x00")):
            raise BoxError("The source filename is not valid for a snapshot copy.")
        try:
            item = self.request("POST", f"/2.0/files/{source_id}/copy", retry=False,
                                params={"fields": "id,type,name,size,extension,file_version,parent"},
                                json={"parent": {"id": destination_id}, "name": name, "version": version_id})
        except BoxError as error:
            if error.status_code == 409:
                raise BoxError("A file with this name already exists in the selected materials category. No file was overwritten.", status_code=409) from None
            if error.ambiguous:
                raise BoxError("Box did not confirm whether the snapshot copy completed. Check the materials folder before trying again; the copy was not retried.", ambiguous=True) from None
            raise
        try:
            target_id = box_id(item["id"])
            if target_id == source_id or str((item.get("parent") or {}).get("id")) != destination_id:
                raise ValueError
            return {**{key: item[key] for key in ("name", "size", "extension", "file_version") if key in item},
                    "id": target_id, "url": f"https://app.box.com/file/{target_id}"}
        except (KeyError, ValueError, TypeError, BoxError):
            raise BoxError("Box returned an incomplete snapshot result. Check the materials folder before trying again; the copy was not retried.", ambiguous=True) from None

    def discover(self, folder_id: str, max_files: int = 25) -> dict:
        queue = deque([box_id(folder_id)])
        visited, files, skipped = set(), [], []
        discovered, truncated = 0, False
        while queue and len(visited) < 100:
            current = queue.popleft()
            if current in visited:
                continue
            visited.add(current)
            folder = self.list_folder(current)
            truncated = truncated or folder["truncated"]
            for item in folder["entries"]:
                if item["type"] == "folder":
                    if item["name"] == "Year in Review outputs":
                        skipped.append({"id": item["id"], "name": item["name"], "reason": "Generated reports are excluded from source materials.", "type": "folder"})
                        continue
                    queue.append(item["id"])
                    continue
                if item["type"] != "file":
                    continue
                discovered += 1
                extension = (item.get("extension") or Path(item["name"]).suffix.lstrip(".")).lower()
                item["extension"] = extension
                reason = None
                if extension not in EXTENSIONS:
                    reason = "Unsupported file format"
                elif item.get("size", 0) > MAX_FILE_BYTES:
                    reason = "File exceeds 12 MB"
                elif len(files) >= max_files:
                    reason = "File limit reached for this run"
                    truncated = True
                if reason:
                    skipped.append({"id": item["id"], "name": item["name"], "reason": reason})
                else:
                    files.append(item)
            if len(files) >= max_files and queue:
                truncated = True
                break
        return {"files": files, "skipped": skipped, "discovered": discovered, "truncated": truncated or bool(queue)}

    def download(self, file: dict) -> bytes:
        file_id = box_id(file["id"])
        version = file.get("file_version") or {}
        url = f"https://api.box.com/2.0/files/{file_id}/content"
        if version.get("id"):
            url += "?version=" + box_id(version["id"])
        for _ in range(6):
            location = urlsplit(url)
            host = location.hostname or ""
            if location.scheme != "https" or not (host == "api.box.com" or host.endswith(".boxcloud.com")):
                raise BoxError("The Box download destination is outside the allowed domains.")
            headers = {"Authorization": "Bearer " + self.token()} if host == "api.box.com" else {}
            try:
                with self.http.stream("GET", url, headers=headers) as response:
                    if response.status_code in {301, 302, 303, 307, 308}:
                        url = response.headers.get("location", "")
                        continue
                    if response.status_code != 200:
                        raise BoxError(f"Cannot read the file (Box HTTP {response.status_code}).")
                    data = bytearray()
                    for chunk in response.iter_bytes():
                        data.extend(chunk)
                        if len(data) > MAX_FILE_BYTES:
                            raise BoxError("The download exceeds 12 MB.")
                    return bytes(data)
            except httpx.HTTPError:
                raise BoxError("The download was interrupted. Please retry.") from None
        raise BoxError("Too many file download redirects.")

    def read_text(self, file: dict) -> str:
        try:
            text = extract_text(self.download(file), file["extension"])
        except BoxError:
            raise
        except Exception:
            raise BoxError("Cannot parse this document. Check its format or save it as plain text.") from None
        if not text.strip():
            raise BoxError("No readable text was found. Scanned documents need OCR first.")
        return text

    def create_folder(self, name: str, parent_id: str = "0") -> dict:
        return self.request("POST", "/2.0/folders", json={"name": name, "parent": {"id": box_id(parent_id)}})

    def upload(self, folder_id: str, name: str, data: bytes, mime_type: str = "application/octet-stream") -> dict:
        if len(data) > MAX_FILE_BYTES:
            raise BoxError("The upload exceeds 12 MB.")
        if not name or len(name) > 200 or any(c in name for c in ("/", "\\", "\r", "\n", "\x00")):
            raise BoxError("The filename is invalid.")
        payload = self.request("POST", "/2.0/files/content", data={"attributes": json.dumps({"name": name, "parent": {"id": box_id(folder_id)}})}, files={"file": (name, data, mime_type)})
        item = payload["entries"][0]
        return {"id": item["id"], "name": item["name"], "url": f"https://app.box.com/file/{item['id']}"}

    def ask(self, file_ids: list[str], prompt: str) -> dict:
        if not 1 <= len(file_ids) <= 25:
            raise BoxError("Box AI requires between 1 and 25 files per batch.")
        return self.request("POST", "/2.0/ai/ask", json={"mode": "single_item_qa" if len(file_ids) == 1 else "multiple_item_qa", "items": [{"id": box_id(i), "type": "file"} for i in file_ids], "include_citations": True, "prompt": prompt[:10000]})

    def upload_version(self, file_id: str, name: str, data: bytes, mime_type: str, etag: str | None = None) -> dict:
        """Replace content as a recoverable Box version, preserving file links."""
        file_id = box_id(file_id)
        if len(data) > MAX_FILE_BYTES or not name or any(c in name for c in ("/", "\\", "\r", "\n", "\x00")):
            raise BoxError("The new file version is too large or has an invalid filename.")
        if etag is not None and not re.fullmatch(r"[A-Za-z0-9._:-]{1,120}", etag):
            raise BoxError("The current Box output version could not be verified. Refresh before saving.")
        try:
            payload = self.request("POST", f"/2.0/files/{file_id}/content", retry=False,
                                   headers={"If-Match": etag} if etag is not None else {},
                                   data={"attributes": json.dumps({"name": name})}, files={"file": (name, data, mime_type)})
            item = payload["entries"][0]
            if str(item["id"]) != file_id or not isinstance(item["name"], str):
                raise ValueError
            return {"id": file_id, "name": item["name"], "url": f"https://app.box.com/file/{file_id}"}
        except BoxError as error:
            if error.status_code == 412:
                raise BoxError("The Box output changed before this save. Its content was not overwritten; check the file and retry.", status_code=412) from None
            if error.ambiguous:
                raise BoxError("Box did not confirm the new file version. Check the existing file before retrying; the upload was not retried.", ambiguous=True) from None
            raise
        except (KeyError, IndexError, TypeError, ValueError):
            raise BoxError("Box did not return a verifiable file version. Check the existing file before retrying; the upload was not retried.", ambiguous=True) from None
