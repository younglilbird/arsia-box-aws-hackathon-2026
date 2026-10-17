"""Downloadable employee reports with compact, linked source references.

Only a small Markdown subset is interpreted: headings, bullets, numbered lists,
and bold text. User/model content never becomes arbitrary HTML or OOXML.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from urllib.parse import urlsplit
from xml.sax.saxutils import escape, quoteattr


_BAD_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
_INLINE = re.compile(r"(\[[A-Za-z][A-Za-z0-9_-]*\]|\*\*[^*\n]+\*\*)")


def _text(value: object) -> str:
    return _BAD_XML.sub("", str(value if value is not None else ""))


def _url(value: object) -> str:
    candidate = _text(value).strip()
    if not candidate or any(c.isspace() for c in candidate) or any(c in candidate for c in '\\<>"'):
        return ""
    try:
        parsed = urlsplit(candidate)
        if parsed.scheme.lower() in {"http", "https"} and parsed.hostname and not parsed.username and not parsed.password:
            return candidate
    except ValueError:
        pass
    return ""


@dataclass(frozen=True)
class Reference:
    id: str
    name: str
    claim: str
    quote: str
    url: str
    source_id: str = ""


def _citation_label(identifier: str) -> str:
    match = re.fullmatch(r"[A-Za-z]+(\d+)", identifier)
    return f"[{int(match.group(1))}]" if match else "[Source]"


def _references(run: dict) -> list[Reference]:
    sources = {str(s.get("id", "")): s for s in run.get("sources", []) if isinstance(s, dict)}
    refs = []
    seen = set()
    for index, evidence in enumerate(run.get("evidence", []), start=1):
        if not isinstance(evidence, dict):
            continue
        identifier = _text(evidence.get("id") or f"E{index}")
        if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_-]*", identifier) or identifier in seen:
            identifier = f"E{index}"
            while identifier in seen:
                identifier += "x"
        seen.add(identifier)
        source = sources.get(str(evidence.get("source_id", "")), {})
        refs.append(Reference(
            identifier,
            _text(evidence.get("source_name") or source.get("name") or "Source"),
            _text(evidence.get("claim", "")),
            _text(evidence.get("quote", "")),
            _url(evidence.get("url")) or _url(source.get("url")),
            _text(evidence.get("source_id", "")),
        ))
    if not refs:
        for index, source in enumerate(sources.values(), start=1):
            refs.append(Reference(f"S{index}", _text(source.get("name") or "Source"), "", "", _url(source.get("url"))))
    return refs


def _cited_source_groups(report: str, refs: list[Reference]) -> list[tuple[list[str], str, str]]:
    """Group cited evidence IDs by source, without repeating their full text."""
    cited = set(re.findall(r"\[([A-Za-z][A-Za-z0-9_-]*)\]", report))
    grouped = {}
    for ref in refs:
        if ref.id not in cited:
            continue
        key = ref.source_id or ref.url or ref.name
        if key not in grouped:
            grouped[key] = ([], ref.name, ref.url)
        grouped[key][0].append(ref.id)
    return list(grouped.values())


def _blocks(markdown: str) -> list[tuple[str, str]]:
    """Parse structure only; raw HTML, tables and embedded links stay text."""
    result = []
    pending = []

    def flush():
        if pending:
            result.append(("p", " ".join(pending)))
            pending.clear()

    for raw in _text(markdown).splitlines():
        line = raw.strip()
        if not line:
            flush()
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        bullet = re.match(r"^[-*+]\s+(.+)$", line)
        numbered = re.match(r"^\d+[.)]\s+(.+)$", line)
        if heading:
            flush()
            result.append((f"h{min(len(heading.group(1)), 3)}", heading.group(2)))
        elif bullet or numbered:
            flush()
            result.append(("bullet" if bullet else "number", (bullet or numbered).group(1)))
        elif re.fullmatch(r"[-*_]{3,}", line):
            flush()
        else:
            pending.append(line)
    flush()
    return result


def _report_body_blocks(run: dict, labels: dict) -> list[tuple[str, str]]:
    """Omit a repeated report title, retaining distinct titles and sections."""
    blocks = _blocks(run["report_markdown"])
    if not blocks or blocks[0][0] != "h1":
        return blocks

    def normalized(value: object) -> str:
        return " ".join(re.findall(r"\w+", _text(value).casefold()))

    heading = " " + normalized(blocks[0][1]) + " "
    title = " " + normalized(labels["title"]) + " "
    if title not in heading:
        return blocks
    remainder = heading.replace(title, " ", 1)
    for detail in (run.get("employee_name"), run.get("year"), "for", "by"):
        phrase = normalized(detail)
        if phrase:
            remainder = remainder.replace(" " + phrase + " ", " ")
    return blocks[1:] if not remainder.strip() else blocks


def _inline_parts(value: str, references: dict[str, Reference]):
    for part in _INLINE.split(value):
        if not part:
            continue
        if part.startswith("[") and part.endswith("]") and part[1:-1] in references:
            yield "reference", _citation_label(part[1:-1]), references[part[1:-1]].url
        elif part.startswith("**") and part.endswith("**"):
            for inner_kind, inner_text, inner_url in _inline_parts(part[2:-2], references):
                yield "bold" if inner_kind == "text" else inner_kind, inner_text, inner_url
        else:
            yield "text", part, ""


def _labels(run: dict):
    confirmed = run.get("reviewer_confirmed") is True
    return {
        "title": "Annual self-review", "sources": "Evidence and sources", "quote": "Source excerpt", "claim": "Supports",
        "state": "Confirmed by employee" if confirmed else "Draft - not yet confirmed by employee",
        "note": "Use these sources to verify the references in the report.", "page": "Page",
    }


def _md_escape(value: str) -> str:
    # Escaping model-supplied Markdown prevents HTML and untrusted links from
    # becoming active when this download is opened in a Markdown viewer.
    value = value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return re.sub(r"([\\`*_{}\[\]()#!|])", r"\\\1", value)


def _md_inline(value: str, refs: dict[str, Reference]) -> str:
    parts = []
    for kind, text, url in _inline_parts(value, refs):
        if kind == "reference" and url:
            parts.append(f"[{_md_escape(text)}](<{url}>)")
        elif kind == "bold":
            parts.append(f"**{_md_escape(text)}**")
        else:
            parts.append(_md_escape(text))
    return "".join(parts)


def _markdown(run: dict, labels: dict, refs: list[Reference]) -> bytes:
    mapping = {ref.id: ref for ref in refs}
    lines = [f"# {labels['title']}", "", f"{_md_escape(_text(run.get('employee_name')))} - {_md_escape(_text(run.get('year')))}", "", f"**{labels['state']}**", ""]
    for kind, value in _blocks(run["report_markdown"]):
        prefix = "#" * (min(int(kind[1:]) + 1, 4)) + " " if kind.startswith("h") else "- " if kind == "bullet" else "1. " if kind == "number" else ""
        lines.extend([prefix + _md_inline(value, mapping), ""])
    if refs:
        lines.extend(["---", "", f"## {labels['sources']}", "", labels["note"], ""])
        for ref in refs:
            label = _md_escape(_citation_label(ref.id))
            link = f"[{label}](<{ref.url}>)" if ref.url else label
            lines.extend([f"{link} **{_md_escape(ref.name)}**", ""])
            if ref.claim:
                lines.extend([f"{labels['claim']}: {_md_escape(ref.claim)}", ""])
            if ref.quote:
                lines.extend([f"{labels['quote']}: {_md_escape(ref.quote)}", ""])
    return "\n".join(lines).encode("utf-8")


def _docx_link(paragraph, text: str, url: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE

    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    properties.append(style)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "176B87")
    properties.append(color)
    run.append(properties)
    content = OxmlElement("w:t")
    content.text = text
    run.append(content)
    link.append(run)
    paragraph._p.append(link)


def _docx_inline(paragraph, value: str, refs: dict[str, Reference]):
    for kind, text, url in _inline_parts(value, refs):
        if kind == "reference" and url:
            _docx_link(paragraph, text, url)
        else:
            run = paragraph.add_run(text)
            run.bold = kind == "bold"


def _docx(run: dict, labels: dict, refs: list[Reference]) -> bytes:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.72)
    section.left_margin = section.right_margin = Inches(0.82)
    section.header_distance = section.footer_distance = Inches(0.30)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number"):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor.from_string("243444")
    normal = document.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.20
    normal.paragraph_format.space_after = Pt(6)
    for name, size in (("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 11)):
        style = document.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
    document.add_heading(labels["title"], 0)
    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.add_run(f"{_text(run.get('employee_name'))} - {_text(run.get('year'))}")
    state = document.add_paragraph(labels["state"])
    state.runs[0].italic = True
    state.runs[0].font.size = Pt(9)
    mapping = {ref.id: ref for ref in refs}
    for kind, value in _report_body_blocks(run, labels):
        style = f"Heading {min(int(kind[1:]), 3)}" if kind.startswith("h") else "List Bullet" if kind == "bullet" else "List Number" if kind == "number" else "Normal"
        paragraph = document.add_paragraph(style=style)
        _docx_inline(paragraph, value, mapping)
    grouped_sources = _cited_source_groups(run["report_markdown"], refs)
    if grouped_sources:
        document.add_heading(labels["sources"], 1)
        for identifiers, name, url in grouped_sources:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.keep_together = True
            citation = ", ".join(_citation_label(identifier) for identifier in identifiers)
            if url:
                _docx_link(paragraph, citation, url)
            else:
                paragraph.add_run(citation)
            paragraph.add_run(" " + name)
            for item in paragraph.runs:
                item.font.size = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run(labels["page"] + " ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for item in footer.runs:
        item.font.size = Pt(8)
    document.core_properties.title = labels["title"]
    document.core_properties.author = _text(run.get("employee_name"))
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _pdf_inline(value: str, mapping: dict[str, Reference]) -> str:
    chunks = []
    for kind, text, url in _inline_parts(value, mapping):
        safe = escape(text)
        if kind == "reference" and url:
            chunks.append(f'<link href={quoteattr(url)} color="#176B87"><u>{safe}</u></link>')
        elif kind == "bold":
            chunks.append(f"<b>{safe}</b>")
        else:
            chunks.append(safe)
    return "".join(chunks)


def _pdf(run: dict, labels: dict, refs: list[Reference]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

    font = "Helvetica"
    base = dict(fontName=font, fontSize=10.5, leading=15.5, textColor=colors.HexColor("#243444"), alignment=TA_LEFT, spaceAfter=7, splitLongWords=True)
    normal = ParagraphStyle("ReportBody", **base)
    styles = {"p": normal, "bullet": ParagraphStyle("ReportBullet", parent=normal, leftIndent=12, firstLineIndent=-9), "number": ParagraphStyle("ReportNumber", parent=normal, leftIndent=12, firstLineIndent=-9)}
    for level, size in ((1, 15), (2, 12.5), (3, 11)):
        styles[f"h{level}"] = ParagraphStyle(f"ReportHeading{level}", parent=normal, fontName="Helvetica-Bold", fontSize=size, leading=size + 5, spaceBefore=11, spaceAfter=5, keepWithNext=True)
    title = ParagraphStyle("ReportTitle", parent=normal, fontName="Helvetica-Bold", fontSize=23, leading=29, spaceAfter=9)
    small = ParagraphStyle("ReportSmall", parent=normal, fontSize=9, leading=13, textColor=colors.HexColor("#5D6B78"))
    output = io.BytesIO()
    document = SimpleDocTemplate(output, pagesize=A4, rightMargin=21 * mm, leftMargin=21 * mm, topMargin=18 * mm, bottomMargin=18 * mm, title=labels["title"], author=_text(run.get("employee_name")))
    story = [Paragraph(escape(labels["title"]), title), Paragraph(escape(f"{_text(run.get('employee_name'))} - {_text(run.get('year'))}"), normal), Paragraph(escape(labels["state"]), small), Spacer(1, 7)]
    mapping = {ref.id: ref for ref in refs}
    number = 0
    for kind, value in _report_body_blocks(run, labels):
        if kind == "number":
            number += 1
        else:
            number = 0
        prefix = "- " if kind == "bullet" else f"{number}. " if kind == "number" else ""
        story.append(Paragraph(prefix + _pdf_inline(value, mapping), styles[kind]))
    grouped_sources = _cited_source_groups(run["report_markdown"], refs)
    if grouped_sources:
        story.append(Paragraph(escape(labels["sources"]), styles["h1"]))
        for identifiers, name, url in grouped_sources:
            identifier = ", ".join(_citation_label(item) for item in identifiers)
            link = f'<link href={quoteattr(url)} color="#176B87"><u>{identifier}</u></link>' if url else identifier
            story.append(Paragraph(link + " " + escape(name), small))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(font, 8)
        canvas.setFillColor(colors.HexColor("#5D6B78"))
        canvas.drawRightString(A4[0] - 21 * mm, 10 * mm, f"{labels['page']} {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def make_export(run: dict, format: str) -> tuple[bytes, str, str]:
    """Return ``(data, download_filename, mime_type)`` for a report run.

    ``format`` accepts ``docx``, ``pdf``, ``md`` or ``markdown``. The report body
    is never shortened to meet a page target; actual pagination depends on its
    content. A compact index of cited sources follows the report body. Full
    evidence text is retained in the Markdown export and the review interface.
    """
    if not isinstance(run, dict) or not isinstance(run.get("report_markdown"), str) or not run["report_markdown"].strip():
        raise ValueError("A nonempty report_markdown is required before exporting.")
    selected = str(format).lower().lstrip(".")
    selected = "md" if selected == "markdown" else selected
    if selected not in {"md", "docx", "pdf"}:
        raise ValueError("Export format must be docx, pdf or markdown.")
    labels = _labels(run)
    refs = _references(run)
    name = re.sub(r"[^\w-]+", "-", _text(run.get("employee_name", "employee")), flags=re.UNICODE).strip("-_")[:65] or "employee"
    year = re.sub(r"[^0-9]", "", _text(run.get("year", "")))[:4] or "year"
    filename = f"annual-review-{year}-{name}.{selected}"
    if selected == "md":
        return _markdown(run, labels, refs), filename, "text/markdown; charset=utf-8"
    if selected == "docx":
        return _docx(run, labels, refs), filename, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    return _pdf(run, labels, refs), filename, "application/pdf"
