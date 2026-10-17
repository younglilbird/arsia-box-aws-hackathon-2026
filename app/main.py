"""Loopback-only, single-workspace annual-review workbench API."""
from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
from contextlib import asynccontextmanager
from copy import deepcopy
from datetime import datetime, timezone
import asyncio
import hashlib
from io import BytesIO
import json
import logging
import math
from pathlib import Path
import re
from threading import RLock
from typing import Literal
from uuid import uuid4

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse, Response
from pydantic import BaseModel, ConfigDict, Field
from starlette.middleware.trustedhost import TrustedHostMiddleware

from app.box_client import BoxClient, BoxError, MAX_FILE_BYTES, box_id
from app.config import ROOT, load_settings
from app.exports import make_export
from app.native import CATEGORIES, NativeBoxAdapter, NativeWorkflowError, _configuration, _review_document, _snapshot
from app.workflow import WorkflowError, run_workflow

ORIGINS = ["http://localhost:5173", "http://127.0.0.1:5173", "http://localhost:8765", "http://127.0.0.1:8765"]


def _now():
    return datetime.now(timezone.utc).isoformat()


def _private_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
    path.parent.chmod(0o700)
    temporary = path.with_suffix(".tmp")
    temporary.touch(mode=0o600, exist_ok=True)
    temporary.chmod(0o600)
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2))
    temporary.replace(path)


def _safe_error(error):
    if isinstance(error, (BoxError, WorkflowError, NativeWorkflowError)):
        return str(error)
    return "The operation could not complete. Check the local AWS sign-in, Box access and network, then retry."


def _file(item):
    return {key: item[key] for key in ("name", "url", "size", "extension", "document_category", "format", "kind", "origin_url") if key in item}


def _save_key(run):
    if not isinstance(run.get("evaluations"), dict):
        return hashlib.sha256((run.get("report_markdown", "") + str(run.get("reviewer_confirmed", False))).encode()).hexdigest()[:10]
    state = {"report_markdown": run.get("report_markdown", ""), "reviewer_confirmed": run.get("reviewer_confirmed", False), "evaluations": run.get("evaluations")}
    return hashlib.sha256(json.dumps(state, sort_keys=True, ensure_ascii=False).encode()).hexdigest()[:10]


def _report_hash(run):
    return hashlib.sha256(run.get("report_markdown", "").encode()).hexdigest()


def _current_box_files(run):
    """A reused Box file exposes only its most recently saved content."""
    current = {}
    for item in run.get("box_files", []):
        if isinstance(item, dict) and item.get("id") is not None:
            current[str(item["id"])] = item
    return list(current.values())


def _public_evaluations(run):
    recorded = run.get("evaluations")
    if not isinstance(recorded, dict):
        return None
    result = {key: deepcopy(recorded[key]) for key in ("rubric_version", "score_label", "overall_delta") if key in recorded}
    fields = {"status", "report_hash", "evaluated_at", "overall_score", "summary", "strengths", "improvements", "missing_information"}
    for name in ("initial", "final"):
        evaluation = recorded.get(name)
        if not isinstance(evaluation, dict):
            continue
        safe = {key: deepcopy(value) for key, value in evaluation.items() if key in fields}
        for collection, allowed in (
            ("dimensions", {"id", "name", "weight", "score", "rationale", "strengths", "improvements", "evidence_ids"}),
            ("leadership_alignment", {"principle", "explanation", "evidence_ids"}),
            ("issues", {"severity", "message", "evidence_ids", "suggestion", "verification_status"}),
        ):
            safe[collection] = [{key: deepcopy(value) for key, value in item.items() if key in allowed} for item in evaluation.get(collection, []) if isinstance(item, dict)]
        result[name] = safe
    final = result.get("final")
    current_hash = hashlib.sha256(run.get("report_markdown", "").encode()).hexdigest()
    if isinstance(final, dict) and final.get("report_hash") != current_hash:
        final["status"] = "stale"
        result["overall_delta"] = None
    return result


def _public(run):
    allowed = {"id", "employee_name", "year", "language", "status", "stage", "progress", "created_at", "updated_at", "completed_at", "events", "report_markdown", "initial_draft", "missing_information", "coverage", "usage", "reviewer_confirmed", "confirmed_at", "error", "save_warnings", "addendum_file", "user_edited", "improvement_error"}
    result = {key: deepcopy(value) for key, value in run.items() if key in allowed}
    result["report_hash"] = _report_hash(run)
    legacy_errors = {
        "The report could not pass source-citation validation after one retry. Retry with a smaller set of documents.": "The generated report did not pass validation. Your materials are unchanged. Try generating it again.",
        "The model returned invalid structured output twice. Retry with fewer or shorter documents.": "The model response did not pass validation. Your materials are unchanged. Try again.",
    }
    if result.get("status") == "failed" and isinstance(result.get("error"), str):
        result["error"] = legacy_errors.get(result["error"], result["error"])
    for event in result.get("events", []):
        if isinstance(event, dict) and isinstance(event.get("message"), str):
            event["message"] = legacy_errors.get(event["message"], event["message"])
    result["sources"] = [_file(item) for item in run.get("sources", [])]
    extracting = run.get("status") in {"queued", "running"} and run.get("stage") in {"discover", "read", "extract"}
    result["evidence"] = [] if extracting else [{key: value for key, value in item.items() if key not in {"source_id"}} for item in run.get("evidence", [])]
    review = deepcopy(run.get("review", {}))
    if isinstance(review.get("box_ai"), dict):
        ai = review["box_ai"]
        ai["file_count"] = len(ai.pop("file_ids", []))
        ai["citations"] = [{key: value for key, value in item.items() if key in {"name", "type"}} for item in ai.get("citations", []) if isinstance(item, dict)]
    result["review"] = review
    if isinstance(result.get("coverage"), dict):
        result["coverage"]["skipped"] = [{key: value for key, value in item.items() if key != "id"} for item in result["coverage"].get("skipped", []) if isinstance(item, dict)]
    result["box_files"] = [_file(item) for item in _current_box_files(run) if item.get("save_key") == _save_key(run)]
    result["saved_file_history"] = [_file(item) for item in run.get("box_files", [])]
    evaluations = _public_evaluations(run)
    if evaluations is not None:
        result["evaluations"] = evaluations
    result["deletion_preview"] = {"tracked_output_count": len(_tracked_outputs(run)), "can_delete": run.get("status") not in {"queued", "running"}}
    if isinstance(run.get("deletion"), dict):
        result["deletion"] = _deletion_summary(run["deletion"])
    if isinstance(run.get("improvement"), dict):
        result["improvement"] = {key: deepcopy(value) for key, value in run["improvement"].items() if key in {"round", "status", "target_score", "started_at", "completed_at", "report_changed", "previous_score", "current_score", "target_reached", "changed_sections"}}
    return result


def _tracked_outputs(run):
    outputs = {}
    for field in ("box_files", "saved_outputs"):
        values = run.get(field, [])
        if isinstance(values, dict):
            values = list(values.values())
        for item in values if isinstance(values, list) else []:
            if isinstance(item, dict) and re.fullmatch(r"[0-9]{1,30}", str(item.get("id", ""))):
                outputs[str(item["id"])] = deepcopy(item)
    return outputs


def _deletion_summary(state):
    statuses = [item.get("status") for item in state.get("files", {}).values() if isinstance(item, dict)]
    return {"status": state.get("status", "partial"), "removed_count": statuses.count("removed"),
            "already_absent_count": statuses.count("absent"), "shared_preserved_count": statuses.count("shared"),
            "remaining_count": sum(status not in {"removed", "absent", "shared"} for status in statuses) + state.get("unresolved_upload_count", 0),
            "warnings": list(state.get("warnings", []))}


class Payload(BaseModel):
    model_config = ConfigDict(extra="forbid")


class RunPayload(Payload):
    notes: str = Field(default="", max_length=8000)


class VersionPayload(Payload):
    expected_report_hash: str = Field(pattern=r"^[a-f0-9]{64}$")


class ReportPayload(VersionPayload):
    report_markdown: str = Field(min_length=1, max_length=60000)
    reviewer_confirmed: Literal[False] = False


class SavePayload(Payload):
    reviewer_confirmed: bool = False


class AnswerPayload(VersionPayload):
    answers: str = Field(min_length=1, max_length=16000)
    target_score: float = Field(default=90, ge=1, le=100, allow_inf_nan=False)


class ImportFilePayload(Payload):
    id: str = Field(pattern=r"^[0-9]{1,30}$")
    version_id: str = Field(pattern=r"^[0-9]{1,30}$")


class ImportPayload(Payload):
    category: Literal["Projects", "Peer Feedback", "Annual Plans"]
    files: list[ImportFilePayload] = Field(min_length=1, max_length=10)


class Workbench:
    def __init__(self, settings, box, native_config=None, directory=None, workflow=run_workflow, executor=None, improvement=None):
        self.settings, self.box, self.native_config = settings, box, native_config
        self.directory = Path(directory or settings.data_dir / "workbench" / "runs")
        self.workflow = workflow
        self.improvement_workflow = improvement
        self.owns_executor = executor is None
        self.executor = executor or ThreadPoolExecutor(max_workers=1, thread_name_prefix="annual-review")
        self.lock, self.operations = RLock(), RLock()
        self.runs = {}

    def config(self):
        try:
            raw = self.native_config
            if raw is None:
                raw = json.loads((ROOT / ".local" / "native-config.json").read_text())
            return _configuration(raw)
        except (OSError, ValueError, NativeWorkflowError):
            raise HTTPException(503, "The fixed Box workspace is not configured. Run the local workspace setup first.") from None

    def load(self):
        self.directory.mkdir(parents=True, exist_ok=True, mode=0o700)
        self.directory.chmod(0o700)
        for path in self.directory.glob("*.json"):
            try:
                run = json.loads(path.read_text())
                if not re.fullmatch(r"[a-f0-9]{32}", run["id"]):
                    continue
                tombstone = self.directory.parent / "deleted-runs" / (run["id"] + ".json")
                if tombstone.is_file() and json.loads(tombstone.read_text()).get("deleted") is True:
                    path.unlink(missing_ok=True)
                    continue
                if run.get("status") in {"queued", "running"}:
                    interrupted_stage = run.get("stage")
                    improvement = run.get("improvement", {})
                    checkpoint_id = improvement.get("checkpoint_id", "") if isinstance(improvement, dict) else ""
                    committed = isinstance(improvement, dict) and (improvement.get("content_committed") is True or (interrupted_stage == "save_draft" and improvement.get("status") == "completed"))
                    if committed:
                        final = run.get("evaluations", {}).get("final", {})
                        valid = final.get("status") == "completed" and final.get("report_hash") == hashlib.sha256(run.get("report_markdown", "").encode()).hexdigest()
                        warning = "Box saving was interrupted after the improved report was evaluated. The current report is retained; retry saving to finish the remaining Box file versions."
                        run.update(status="completed" if valid else "failed", stage="completed" if valid else "failed", progress=100 if valid else run.get("progress", 97), updated_at=_now(), improvement_error=warning)
                        run["save_warnings"] = list(dict.fromkeys(run.get("save_warnings", []) + [warning]))
                        if not valid:
                            run["error"] = "The interrupted candidate's evaluation could not be verified. Its stored text has been retained; check local recovery before continuing."
                    elif re.fullmatch(r"[a-f0-9]{32}", checkpoint_id):
                        checkpoint = self.directory.parent / "improvement-checkpoints" / (run["id"] + "-" + checkpoint_id + ".json")
                        try:
                            previous = json.loads(checkpoint.read_text())["run"]
                            if not isinstance(previous, dict) or previous.get("id") != run["id"]:
                                raise ValueError
                            run = self.restored_improvement(previous, run, "The local server restarted during improvement. Your previous report has been restored.")
                        except (OSError, ValueError, KeyError, TypeError):
                            run.update(status="failed", stage="failed", updated_at=_now(), error="The improvement was interrupted and its private checkpoint could not be read. The stored report is retained; check local storage before retrying.")
                    else:
                        run.update(status="failed", stage="failed", updated_at=_now(), error="The local server restarted during this run. Start a new run.")
                    run["interrupted_stage"] = interrupted_stage
                    self.persist(run)
                self.runs[run["id"]] = run
            except (OSError, ValueError, KeyError, TypeError):
                continue

    def shutdown(self):
        if not self.owns_executor:
            return
        with self.lock:
            active = any(run.get("status") in {"queued", "running"} for run in self.runs.values())
        if active:
            logging.getLogger(__name__).warning("Waiting for active generation to finish before shutting down.")
        self.executor.shutdown(wait=True, cancel_futures=False)

    def persist(self, run):
        self.directory.mkdir(parents=True, exist_ok=True, mode=0o700)
        path = self.directory / (run["id"] + ".json")
        temp = path.with_suffix(".tmp")
        temp.touch(mode=0o600, exist_ok=True)
        temp.chmod(0o600)
        temp.write_text(json.dumps(run, ensure_ascii=False, indent=2))
        temp.replace(path)

    def get(self, run_id):
        if not re.fullmatch(r"[a-f0-9]{32}", run_id):
            raise HTTPException(404, "Run not found.")
        with self.lock:
            if run_id not in self.runs:
                raise HTTPException(404, "Run not found.")
            return deepcopy(self.runs[run_id])

    def update(self, run_id, **fields):
        with self.lock:
            run = self.runs[run_id]
            message = fields.pop("message", None)
            if "sources" in fields:
                fields["sources"] = [{key: value for key, value in source.items() if key != "text"} for source in fields["sources"]]
            run.update(deepcopy(fields))
            run["updated_at"] = _now()
            if message:
                run["events"].append({"at": _now(), "stage": run.get("stage"), "message": message})
            self.persist(run)

    def ensure_idle(self):
        with self.lock:
            if any(item.get("status") in {"queued", "running"} for item in self.runs.values()):
                raise HTTPException(409, "A generation is already running. Wait for it to complete.")

    @staticmethod
    def check_report_version(run, expected_hash):
        if _report_hash(run) != expected_hash:
            raise HTTPException(409, "This report changed in another tab. Reload the current report before applying your edits, confirming or improving it. Your submitted changes were not applied.")

    def check_identity(self):
        identity = self.box.identity()
        if self.settings.box_user_id and str(identity.get("id")) != str(self.settings.box_user_id):
            raise BoxError("Box is signed in as a different user from the configured workspace. Check the local Box user configuration.")
        return identity

    def check_import_identity(self):
        if not self.settings.box_user_id:
            raise HTTPException(503, "Configure the human Box user before browsing or importing files.")
        return self.check_identity()

    def import_manifest(self):
        path = self.directory.parent / "imports.json"
        try:
            if not path.exists():
                return []
            value = json.loads(path.read_text())
            if not isinstance(value, dict) or not isinstance(value.get("imports"), list):
                raise ValueError
            if not all(isinstance(item, dict) for item in value["imports"]):
                raise ValueError
            return value["imports"]
        except (OSError, ValueError):
            raise HTTPException(503, "The private import history could not be read. Check local storage before importing again.") from None

    def save_import_manifest(self, records):
        path = self.directory.parent / "imports.json"
        path.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
        path.parent.chmod(0o700)
        temporary = path.with_suffix(".tmp")
        temporary.touch(mode=0o600, exist_ok=True)
        temporary.chmod(0o600)
        temporary.write_text(json.dumps({"version": 1, "imports": records}, ensure_ascii=False, indent=2))
        temporary.replace(path)

    @staticmethod
    def import_exclusions(config):
        return {config["root_folder_id"], config["reports_folder_id"], *config["category_folders"].values()}

    def check_external_source(self, item, config):
        path = item.get("path_collection")
        if not isinstance(path, dict) or not isinstance(path.get("entries"), list):
            raise HTTPException(400, "Box could not verify this item's source location. Choose another file or refresh the folder.")
        ancestors = {box_id(parent.get("id", "")) for parent in path["entries"] if isinstance(parent, dict)}
        if ancestors.intersection(self.import_exclusions(config)):
            raise HTTPException(400, "Files already inside this review workspace cannot be imported again. Generated reports are excluded from source materials.")

    def browse_box(self, folder_id):
        self.check_import_identity()
        config, folder_id = self.config(), box_id(folder_id)
        excluded = self.import_exclusions(config)
        if folder_id in excluded:
            raise HTTPException(400, "Choose source files outside the current review workspace.")
        folder = self.box.list_folder(folder_id)
        if folder_id != "0":
            self.check_external_source(folder, config)
        entries = []
        for item in folder.get("entries", []):
            identifier = box_id(item.get("id", ""))
            kind = item.get("type")
            if identifier in excluded or kind not in {"file", "folder"}:
                continue
            row = {"id": identifier, "type": kind, "name": str(item.get("name", "Untitled")),
                   "url": f"https://app.box.com/{'folder' if kind == 'folder' else 'file'}/{identifier}"}
            if kind == "file":
                extension = str(item.get("extension") or Path(row["name"]).suffix.lstrip(".")).lower()
                version = (item.get("file_version") or {}).get("id")
                size = item.get("size")
                if extension not in {"docx", "pdf"} or not isinstance(size, int) or not 0 < size <= MAX_FILE_BYTES or not re.fullmatch(r"[0-9]{1,30}", str(version or "")):
                    continue
                row.update(size=size, extension=extension, version_id=str(version))
            entries.append(row)
        parent = folder.get("parent")
        safe_parent = {"id": box_id(parent["id"]), "name": str(parent.get("name", "Box"))} if isinstance(parent, dict) and parent.get("id") is not None else None
        return {"id": folder_id, "name": str(folder.get("name", "Box")), "parent": safe_parent,
                "entries": entries, "truncated": bool(folder.get("truncated"))}

    def import_files(self, payload):
        with self.operations:
            self.ensure_idle()
            self.check_import_identity()
            config = self.config()
            if len({item.id for item in payload.files}) != len(payload.files):
                raise HTTPException(400, "Select each Box file only once per import.")
            records = self.import_manifest()
            prepared = []
            for selection in payload.files:
                info = self.box.file_info(selection.id)
                if str(info.get("id")) != selection.id or info.get("type") != "file":
                    raise HTTPException(400, "Select a supported Box file, not a folder.")
                self.check_external_source(info, config)
                name, size = str(info.get("name", "")), info.get("size")
                extension = str(info.get("extension") or Path(name).suffix.lstrip(".")).lower()
                if extension not in {"docx", "pdf"} or not isinstance(size, int) or not 0 < size <= MAX_FILE_BYTES or not name or len(name) > 255 or any(char in name for char in ("/", "\\", "\r", "\n", "\x00")):
                    raise HTTPException(400, "Snapshot imports support nonempty Word/PDF files up to 12 MB only.")
                if str((info.get("file_version") or {}).get("id")) != selection.version_id:
                    raise HTTPException(409, "A selected Box source changed after browsing. Refresh the folder and select its current version before importing.")
                prepared.append({"id": selection.id, "version_id": selection.version_id, "name": name, "size": size, "extension": extension})
            # Confirm local provenance storage is writable before any Box copy.
            try:
                self.save_import_manifest(records)
            except OSError:
                raise HTTPException(503, "The private import history could not be saved. No Box files were copied.") from None
            copied, warnings = [], []
            for source in prepared:
                try:
                    target = self.box.copy_file(source["id"], config["category_folders"][payload.category], source["name"], source["version_id"])
                except BoxError as error:
                    warnings.append(source["name"] + ": " + _safe_error(error))
                    if getattr(error, "ambiguous", False):
                        warnings.append("Remaining selections were not copied because the last outcome is uncertain. Check the materials folder before another import.")
                        break
                    continue
                origin_url = "https://app.box.com/file/" + source["id"]
                copied.append(_file({**source, **target, "document_category": payload.category, "origin_url": origin_url}))
                records.append({"source_id": source["id"], "source_version_id": source["version_id"], "source_name": source["name"],
                                "target_id": box_id(target["id"]), "target_url": target["url"], "category": payload.category,
                                "origin_url": origin_url, "imported_at": _now(), "mode": "snapshot"})
                try:
                    self.save_import_manifest(records)
                except OSError:
                    warnings.append("A snapshot was copied, but its private origin history could not be saved. Check the materials folder before repeating the import. Remaining selections were not copied.")
                    break
            return {"files": copied, "warnings": warnings}

    def delete_run(self, run_id):
        if not re.fullmatch(r"[a-f0-9]{32}", run_id):
            raise HTTPException(404, "Run not found.")
        with self.operations:
            tombstone = self.directory.parent / "deleted-runs" / (run_id + ".json")
            with self.lock:
                exists = run_id in self.runs
            if not exists:
                if tombstone.is_file():
                    try:
                        result = json.loads(tombstone.read_text())
                        if result.get("id") == run_id and result.get("deleted") is True:
                            return {key: deepcopy(result[key]) for key in ("id", "deleted", "status", "removed_count", "already_absent_count", "shared_preserved_count", "remaining_count", "warnings")}
                    except (OSError, ValueError, KeyError):
                        pass
                raise HTTPException(404, "Run not found.")
            run = self.get(run_id)
            if run.get("status") in {"queued", "running"}:
                raise HTTPException(409, "A running generation cannot be deleted. Wait for it to finish.")
            outputs, pending = _tracked_outputs(run), deepcopy(run.get("pending_uploads", {}))
            identity, config = (self.check_import_identity(), self.config()) if outputs or pending else (None, None)
            files, recovery_warnings = list(run.get("box_files", [])), []
            for key, intent in list(pending.items()):
                try:
                    recovered = self.recover_upload(run, intent, identity, config)
                    if recovered:
                        files.append({**recovered, **{field: intent[field] for field in ("kind", "format", "save_key")}})
                        pending.pop(key)
                except BoxError:
                    recovery_warnings.append("A previous upload could not be verified safely. Check Reports and retry deletion; no unverified file was removed.")
            if run.get("pending_uploads"):
                self.update(run_id, box_files=files, pending_uploads=pending)
                run, outputs = self.get(run_id), _tracked_outputs({"box_files": files, "saved_outputs": run.get("saved_outputs", [])})
            if pending:
                state = {"status": "partial", "unresolved_upload_count": len(pending), "requested_at": _now(),
                         "files": {identifier: {"name": item.get("name", "Report output"), "status": "pending"} for identifier, item in outputs.items()},
                         "warnings": list(dict.fromkeys(recovery_warnings + ["A previous Box upload has an uncertain outcome. The draft history was retained until every generated file can be verified."]))}
                self.update(run_id, deletion=state)
                return {"id": run_id, "deleted": False, **_deletion_summary(state), "run": _public(self.get(run_id))}
            def confirmed_trashed(identifier):
                try:
                    info = self.box.trashed_file_info(identifier)
                    return str(info.get("id")) == identifier and info.get("type") == "file" and info.get("item_status") == "trashed" and str((info.get("owned_by") or {}).get("id")) == str(identity.get("id"))
                except Exception:
                    return False
            unavailable = "A tracked Box output is unavailable, and its removal could not be verified. The draft history was retained; check Box access before retrying."
            with self.lock:
                other_runs = [deepcopy(value) for key, value in self.runs.items() if key != run_id]
            shared = {identifier for other in other_runs for identifier in _tracked_outputs(other)}
            source_ids = {str(source.get("id")) for item in [run, *other_runs] for source in item.get("sources", []) if isinstance(source, dict)}
            previous = run.get("deletion") if isinstance(run.get("deletion"), dict) else {}
            state = {"status": "pending", "requested_at": previous.get("requested_at", _now()), "files": deepcopy(previous.get("files", {})), "warnings": []}
            for identifier, output in outputs.items():
                state["files"].setdefault(identifier, {"name": str(output.get("name", "Report output")), "status": "pending"})
            self.update(run_id, deletion=state)
            for identifier, output in outputs.items():
                item = state["files"][identifier]
                if item.get("status") in {"removed", "absent"}:
                    continue
                if identifier in source_ids:
                    item.update(status="blocked", message="A tracked item is also a source document and was not deleted.")
                elif identifier in shared:
                    item.update(status="shared", message="An output is referenced by another draft and was preserved.")
                elif output.get("kind") not in {None, "report", "review_notes"}:
                    item.update(status="blocked", message="A tracked item is not a recognized report output and was not deleted.")
                else:
                    try:
                        info = self.box.file_info(identifier)
                        if info.get("item_status") == "trashed" and str(info.get("id")) == identifier and info.get("type") == "file" and str((info.get("owned_by") or {}).get("id")) == str(identity.get("id")):
                            item.update(status="absent", message="A tracked output was already removed.")
                        elif info.get("item_status") != "active":
                            item.update(status="failed", message=unavailable)
                        elif str(info.get("id")) != identifier or info.get("type") != "file" or str((info.get("parent") or {}).get("id")) != config["reports_folder_id"] or str((info.get("owned_by") or {}).get("id")) != str(identity.get("id")):
                            item.update(status="blocked", message="An output's Reports location or Box owner could not be verified; it was not deleted.")
                        else:
                            result = self.box.trash_file(identifier, str(info.get("etag", "")))
                            if result.get("needs_verification"):
                                verified = confirmed_trashed(identifier)
                                item.update(status="absent" if verified else "failed", message="" if verified else unavailable)
                            else:
                                item.update(status="absent" if result.get("already_absent") else "removed", message="")
                    except BoxError as error:
                        if getattr(error, "status_code", None) == 404:
                            verified = confirmed_trashed(identifier)
                            item.update(status="absent" if verified else "failed", message="A tracked output was verified in Box Trash." if verified else unavailable)
                        elif getattr(error, "ambiguous", False) and confirmed_trashed(identifier):
                            item.update(status="absent", message="A tracked output was verified in Box Trash after an uncertain deletion response.")
                        else:
                            item.update(status="failed", message="A Box output could not be removed. The draft history was retained; retry deletion after checking Box access and the file.")
                    except Exception:
                        item.update(status="failed", message="An output could not be verified or removed. The draft history was retained for a safe retry.")
                state["warnings"] = list(dict.fromkeys(value["message"] for value in state["files"].values() if value.get("message")))
                self.update(run_id, deletion=state)
            summary = _deletion_summary(state)
            if summary["remaining_count"]:
                state["status"] = "partial"
                self.update(run_id, deletion=state)
                return {"id": run_id, "deleted": False, **_deletion_summary(state), "run": _public(self.get(run_id))}
            result = {"id": run_id, "deleted": True, **{**summary, "status": "deleted"}}
            tombstone.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
            tombstone.parent.chmod(0o700)
            temporary = tombstone.with_suffix(".tmp")
            temporary.touch(mode=0o600, exist_ok=True)
            temporary.chmod(0o600)
            temporary.write_text(json.dumps({**result, "deleted_at": _now()}, ensure_ascii=False, indent=2))
            temporary.replace(tombstone)
            try:
                (self.directory / (run_id + ".json")).unlink(missing_ok=True)
            except OSError:
                tombstone.unlink(missing_ok=True)
                state["status"] = "partial"
                state["warnings"].append("Box output removal completed, but local history could not be removed. Retry deletion to finish.")
                self.update(run_id, deletion=state)
                return {"id": run_id, "deleted": False, **_deletion_summary(state), "run": _public(self.get(run_id))}
            with self.lock:
                self.runs.pop(run_id, None)
            return result

    def delete_material(self, file_id):
        identifier = box_id(file_id)
        with self.operations:
            self.ensure_idle()
            identity, config = self.check_import_identity(), self.config()
            with self.lock:
                generated_ids = {value for run in self.runs.values() for value in _tracked_outputs(run)}
            origin_ids = {str(item.get("source_id")) for item in self.import_manifest()}
            if identifier in generated_ids or identifier in origin_ids:
                raise HTTPException(400, "Generated reports and original imported files cannot be deleted through Materials.")
            record_path = self.directory.parent / "material-deletions" / (identifier + ".json")
            previous = None
            if record_path.is_file():
                try:
                    previous = json.loads(record_path.read_text())
                    if not isinstance(previous, dict):
                        raise ValueError
                except (OSError, ValueError):
                    raise HTTPException(503, "The private removal record could not be read. Check local storage before trying again.") from None
                if previous.get("user_id") != str(identity["id"]) or previous.get("workspace_id") != config["root_folder_id"] or previous.get("category") not in config["category_folders"] or previous.get("category_folder_id") != config["category_folders"].get(previous.get("category")):
                    raise HTTPException(409, "This removal record belongs to a different materials workspace.")
            snapshot = _snapshot(self.box, config)
            source = next((item for item in snapshot["files"] if str(item.get("id")) == identifier), None)
            if source is None and previous is None:
                raise HTTPException(404, "This file is not in the current Materials folders. Refresh Materials before deleting a file.")
            category = source["document_category"] if source else previous["category"]
            category_folder = config["category_folders"][category]
            record = {"material_id": identifier, "user_id": str(identity["id"]), "workspace_id": config["root_folder_id"],
                      "category": category, "category_folder_id": category_folder, "status": "pending", "updated_at": _now()}

            def persist_record(status):
                record.update(status=status, updated_at=_now())
                record_path.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
                record_path.parent.chmod(0o700)
                temporary = record_path.with_suffix(".tmp")
                temporary.touch(mode=0o600, exist_ok=True)
                temporary.chmod(0o600)
                temporary.write_text(json.dumps(record, ensure_ascii=False, indent=2))
                temporary.replace(record_path)

            def confirmed_trashed():
                try:
                    info = self.box.trashed_file_info(identifier)
                    return str(info.get("id")) == identifier and info.get("type") == "file" and info.get("item_status") == "trashed" and str((info.get("owned_by") or {}).get("id")) == str(identity["id"])
                except Exception:
                    return False

            def unavailable():
                persist_record("unavailable")
                return {"deleted": False, "status": "unavailable", "warnings": ["The material's removal could not be verified in Box. It was not marked as deleted; check Box access and retry."]}

            if source is None:
                if confirmed_trashed():
                    persist_record("deleted")
                    return {"deleted": True, "status": "deleted", "warnings": []}
                return unavailable()
            try:
                info = self.box.file_info(identifier)
            except BoxError as error:
                if error.status_code == 404 and confirmed_trashed():
                    persist_record("deleted")
                    return {"deleted": True, "status": "deleted", "warnings": []}
                return unavailable()
            if info.get("item_status") == "trashed" and confirmed_trashed():
                persist_record("deleted")
                return {"deleted": True, "status": "deleted", "warnings": []}
            ancestors = {str(item.get("id")) for item in (info.get("path_collection") or {}).get("entries", []) if isinstance(item, dict)}
            parent_id = str((info.get("parent") or {}).get("id"))
            if str(info.get("id")) != identifier or info.get("type") != "file" or info.get("item_status") != "active" or str((info.get("owned_by") or {}).get("id")) != str(identity["id"]) or category_folder not in ancestors or parent_id not in ancestors or config["reports_folder_id"] in ancestors:
                raise HTTPException(409, "The material's owner or folder changed. Refresh Materials before trying to delete it.")
            record.update(parent_id=parent_id, version_id=str((info.get("file_version") or {}).get("id", "")))
            persist_record("pending")
            try:
                removed = self.box.trash_file(identifier, str(info.get("etag", "")))
                if removed.get("needs_verification") and not confirmed_trashed():
                    return unavailable()
            except BoxError as error:
                if (error.ambiguous or error.status_code == 404) and confirmed_trashed():
                    pass
                else:
                    return unavailable()
            persist_record("deleted")
            return {"deleted": True, "status": "deleted", "warnings": []}

    def start(self, notes="", addendum_file=None):
        with self.operations:
            self.ensure_idle()
            config = self.config()
            self.check_identity()
            snapshot = _snapshot(self.box, config)
            if not snapshot["files"]:
                raise HTTPException(400, "Upload Word or text-based PDF documents before starting a run.")
            run_id = uuid4().hex
            run = {"id": run_id, "employee_name": config["employee_name"], "year": config["year"], "language": "en", "status": "queued", "stage": "discover", "progress": 0, "created_at": _now(), "updated_at": _now(), "events": [], "sources": [], "evidence": [], "review": {}, "report_markdown": "", "initial_draft": "", "missing_information": [], "reviewer_confirmed": False, "box_files": [], "fingerprint": snapshot["fingerprint"], "request_notes": notes}
            if addendum_file:
                run["addendum_file"] = _file(addendum_file)
            with self.lock:
                self.runs[run_id] = run
                self.persist(run)
            self.executor.submit(self.execute, run_id, config, snapshot, notes)
            return _public(self.get(run_id))

    @staticmethod
    def restored_improvement(previous, current, message):
        restored = deepcopy(previous)
        restored.update(status="completed", stage="completed", progress=100, updated_at=_now(), improvement_error=message)
        for field in ("request_notes", "improvement_answers", "pending_addendum", "improvement_input_fingerprint", "addendum_upload_intent"):
            if field in current:
                restored[field] = deepcopy(current[field])
        if current.get("improvement_input_fingerprint"):
            restored["fingerprint"] = current["improvement_input_fingerprint"]
        improvement = deepcopy(current.get("improvement", {}))
        improvement.update(status="failed", completed_at=_now(), report_changed=False)
        restored["improvement"] = improvement
        restored["events"] = deepcopy(current.get("events", [])) + [{"at": _now(), "stage": "completed", "message": message}]
        return restored

    def improve(self, run_id, answers, target_score, expected_report_hash):
        if not answers.strip():
            raise HTTPException(400, "Enter additional context before improving this report.")
        with self.operations:
            self.ensure_idle()
            previous, config = self.get(run_id), self.config()
            if previous.get("status") != "completed" or previous.get("deletion"):
                raise HTTPException(409, "Complete this draft before improving it.")
            self.check_report_version(previous, expected_report_hash)
            self.check_import_identity()
            snapshot = _snapshot(self.box, config)
            if snapshot["fingerprint"] != previous.get("fingerprint"):
                self.check_versions(previous, config)
            if snapshot.get("supported_count", len(snapshot["files"])) >= 25 or snapshot.get("truncated"):
                raise HTTPException(400, "The materials workspace is at its input limit. Make room for an employee-answer document before improving this draft.")
            attempt = uuid4().hex
            checkpoint = self.directory.parent / "improvement-checkpoints" / (run_id + "-" + attempt + ".json")
            _private_json(checkpoint, {"run": previous, "saved_at": _now()})
            round_number = 1 + sum(item.get("status") == "completed" for item in previous.get("revision_audit", []))
            improvement = {"round": round_number, "status": "running", "target_score": target_score, "started_at": _now(), "checkpoint_id": attempt}
            answer_history = deepcopy(previous.get("improvement_answers", []))
            answer_history.append({"round": round_number, "answers": answers, "target_score": target_score, "at": _now()})
            request_notes = previous.get("request_notes", "") + "\n\nSelf-reported employee answers, not independently verified:\n" + answers
            self.update(run_id, status="running", stage="improve", progress=5, improvement=improvement, improvement_error=None,
                        request_notes=request_notes, improvement_answers=answer_history,
                        improvement_input_fingerprint=None, pending_addendum=None,
                        message="Improving this draft with employee-supplied context; one revision and one fresh evaluation.")
            try:
                self.executor.submit(self.execute_improvement, run_id, config, snapshot, previous, answers, target_score)
            except Exception:
                with self.lock:
                    restored = self.restored_improvement(previous, self.runs[run_id], "The improvement could not start. Your previous report is still available; try again.")
                    self.runs[run_id] = restored
                    self.persist(restored)
                raise HTTPException(503, "The improvement could not start. Your previous report is still available; try again.") from None
            return _public(self.get(run_id))

    def execute_improvement(self, run_id, config, before_snapshot, previous, answers, target_score):
        committed = False
        try:
            from docx import Document
            from app.box_client import extract_text
            identity = self.check_import_identity()
            if _snapshot(self.box, config)["fingerprint"] != before_snapshot["fingerprint"]:
                raise HTTPException(409, "Source documents changed before improvement. Your previous report has been retained; refresh the materials.")
            document = Document()
            document.add_heading("Employee-supplied annual-review addendum", 0)
            document.add_paragraph("Source classification: self-reported employee context. These answers are not independent measurements or verified team outcomes; retain attribution and verify quantitative claims against project evidence.")
            document.add_paragraph(f"Employee: {config['employee_name']}")
            document.add_paragraph(f"Review year: {config['year']}")
            document.add_paragraph("Date: " + _now())
            document.add_heading("Questions being addressed", level=1)
            for question in previous.get("missing_information", []):
                document.add_paragraph(str(question))
            document.add_heading("Employee answers", level=1)
            for paragraph in answers.splitlines():
                document.add_paragraph(re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", paragraph))
            buffer = BytesIO()
            document.save(buffer)
            name = "Employee-addendum-" + datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S-") + uuid4().hex[:6] + ".docx"
            intent = {"name": name, "folder_id": config["category_folders"]["Annual Plans"], "owner_id": str(identity["id"]),
                      "kind": "employee_addendum", "format": "docx", "content_sha1": [hashlib.sha1(buffer.getvalue()).hexdigest()], "status": "pending"}
            self.update(run_id, addendum_upload_intent=intent)
            try:
                uploaded = self.box.upload(intent["folder_id"], name, buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
            except BoxError as error:
                if not error.ambiguous:
                    self.update(run_id, addendum_upload_intent=None)
                    raise
                self.update(run_id, addendum_upload_intent={**intent, "status": "uncertain"})
                try:
                    uploaded = self.recover_upload(self.get(run_id), intent, identity, config, material=True)
                except BoxError:
                    uploaded = None
                if uploaded is None:
                    raise WorkflowError("The employee-answer upload could not be verified.", "addendum_upload_uncertain") from None
            self.update(run_id, pending_addendum=uploaded, addendum_upload_intent=None, message="Employee answers were saved to Box as self-reported context.")
            snapshot = _snapshot(self.box, config)
            def source_signature(item):
                return [item.get("file_version"), item.get("name"), item.get("size"), item.get("document_category")]
            old_versions = {str(item["id"]): source_signature(item) for item in before_snapshot["files"]}
            new_versions = {str(item["id"]): source_signature(item) for item in snapshot["files"] if str(item["id"]) != str(uploaded["id"])}
            if old_versions != new_versions or snapshot.get("truncated"):
                raise HTTPException(409, "Source documents changed or the input limit was reached during improvement. Your previous report has been retained.")
            source = next((item for item in snapshot["files"] if str(item["id"]) == str(uploaded["id"])), None)
            if source is None:
                raise WorkflowError("The saved answers were not available in the materials snapshot.")
            self.update(run_id, improvement_input_fingerprint=snapshot["fingerprint"])
            candidate = deepcopy(previous)
            candidate["improvement_context"] = {"new_source": {**source, "text": extract_text(buffer.getvalue(), "docx")},
                                                "answered_questions": list(previous.get("missing_information", [])), "snapshot_fingerprint": snapshot["fingerprint"]}
            helper = self.improvement_workflow
            if helper is None:
                from app.improvement import improve_review
                helper = improve_review
            result = helper(NativeBoxAdapter(self.box, config["root_folder_id"], snapshot), self.settings, candidate, answers, target_score,
                            lambda **fields: self.update(run_id, **fields))
            if not isinstance(result, dict) or not isinstance(result.get("report_markdown"), str) or not result["report_markdown"].strip():
                raise WorkflowError("The improvement response was incomplete. Your previous report is still available.")
            latest = result.get("evaluations", {}).get("final", {})
            score = latest.get("overall_score")
            expected_hash = hashlib.sha256(result["report_markdown"].encode()).hexdigest()
            if latest.get("status") != "completed" or latest.get("report_hash") != expected_hash or isinstance(score, bool) or not isinstance(score, (int, float)) or not math.isfinite(score) or not 0 <= score <= 100:
                raise WorkflowError("The improved report did not have a valid fresh evaluation. Your previous report is still available.", "improvement_evaluation")
            evidence = result.get("evidence", [])
            if not isinstance(evidence, list) or len(evidence) > 20 or any(not isinstance(item, dict) for item in evidence):
                raise WorkflowError("The improved evidence did not pass validation.", "improvement_evidence")
            evidence_ids = {item.get("id") for item in evidence}
            if len(evidence_ids) != len(evidence) or not set(re.findall(r"\[(E[0-9]+)\]", result["report_markdown"])) <= evidence_ids:
                raise WorkflowError("The improved report did not pass evidence validation.", "improvement_evidence")
            if _snapshot(self.box, config)["fingerprint"] != snapshot["fingerprint"]:
                raise HTTPException(409, "Source documents changed during improvement. Your previous report has been retained.")
            changed = result["report_markdown"] != previous.get("report_markdown")
            evaluations = deepcopy(result.get("evaluations", {}))
            old_evaluations = previous.get("evaluations", {})
            if isinstance(old_evaluations.get("initial"), dict):
                evaluations["initial"] = deepcopy(old_evaluations["initial"])
                if isinstance(evaluations.get("final", {}).get("overall_score"), (int, float)):
                    evaluations["overall_delta"] = round(evaluations["final"]["overall_score"] - old_evaluations["initial"]["overall_score"], 2)
            improvement = deepcopy(self.get(run_id)["improvement"])
            improvement.update(status="completed", completed_at=_now(), report_changed=changed,
                               content_committed=True,
                               previous_score=old_evaluations.get("final", {}).get("overall_score"), current_score=evaluations.get("final", {}).get("overall_score"),
                               target_reached=score >= target_score, changed_sections=list(result.get("improvement", {}).get("changed_sections", [])))
            audit = deepcopy(previous.get("revision_audit", []))
            audit.append({**improvement, "previous_evaluations": deepcopy(old_evaluations), "new_evaluations": deepcopy(evaluations),
                          "previous_report_hash": hashlib.sha256(previous.get("report_markdown", "").encode()).hexdigest(),
                          "report_hash": hashlib.sha256(result["report_markdown"].encode()).hexdigest(), "addendum_source_id": str(uploaded["id"])})
            slots = deepcopy(previous.get("output_slots", {}))
            for item in previous.get("box_files", []):
                if item.get("kind") in {"report", "review_notes"} and item.get("format") in {"docx", "pdf"}:
                    slots[item["kind"] + ":" + item["format"]] = str(item["id"])
            safe_result = {key: value for key, value in result.items() if key in {"sources", "evidence", "report_markdown", "review", "missing_information", "coverage", "usage", "improvement_history"}}
            self.update(run_id, **safe_result, evaluations=evaluations, initial_draft=previous.get("initial_draft", ""), fingerprint=snapshot["fingerprint"],
                        reviewer_confirmed=False if changed else previous.get("reviewer_confirmed", False), confirmed_at=None if changed else previous.get("confirmed_at"),
                        improvement=improvement, revision_audit=audit, output_slots=slots, reuse_output_ids=True, addendum_file=_file(uploaded),
                        stage="save_draft", progress=97, message="Saving the improved report and evaluation as new versions of the same Box files.")
            committed = True
            self.publish(run_id)
            self.update(run_id, status="completed", stage="completed", progress=100, completed_at=_now(), improvement_error=None,
                        message="Improvement round complete. Review the current evaluation and remaining questions before another manual round.")
        except Exception as error:
            with self.lock:
                current = deepcopy(self.runs[run_id])
                if committed:
                    self.update(run_id, status="completed", stage="completed", progress=100,
                                improvement_error="The improved report is ready locally, but Box saving could not finish. Check the save warnings and retry saving.",
                                save_warnings=["The improved report could not be fully saved to Box. Its local report and fresh evaluation are retained."])
                    return
                answer_status = "Your answers were saved as self-reported material; they have not been applied to the restored report." if isinstance(current.get("pending_addendum"), dict) else "Your answers remain in private local history; no report change was applied."
                if isinstance(current.get("addendum_upload_intent"), dict):
                    answer_status = "Box could not confirm whether the employee-answer document was saved. Check Materials for " + current["addendum_upload_intent"]["name"] + "; remove that unapplied answer copy, if present, before retrying."
                restored = self.restored_improvement(previous, current, "This improvement round could not complete. Your previous report has been restored. " + answer_status)
                if isinstance(error, WorkflowError):
                    restored["internal_diagnostics"] = {"code": str(error.code)[:80], "validation_reason": error.validation_reason[:500] if isinstance(error.validation_reason, str) else None}
                self.runs[run_id] = restored
                self.persist(restored)

    def check_versions(self, run, config):
        if _snapshot(self.box, config)["fingerprint"] != run.get("fingerprint"):
            if run.get("reviewer_confirmed"):
                self.update(run["id"], reviewer_confirmed=False, confirmed_at=None, message="Source documents changed. Confirmation has been reset; generate a fresh draft.")
            raise HTTPException(409, "Source documents changed after this run started. Generate a fresh draft before publishing or confirming it.")

    def recover_upload(self, run, intent, identity, config, *, material=False):
        """Adopt an uncertain upload only after verifying its exact intended bytes."""
        folder_id = config["category_folders"]["Annual Plans"] if material else config["reports_folder_id"]
        kinds, formats = ({"employee_addendum"}, {"docx"}) if material else ({"report", "review_notes"}, {"docx", "pdf"})
        if intent.get("folder_id") != folder_id or intent.get("owner_id") != str(identity.get("id")) or intent.get("kind") not in kinds or intent.get("format") not in formats:
            raise BoxError("The pending upload belongs to a different output workspace. Check local recovery before saving again.")
        folder = self.box.list_folder(folder_id)
        if folder.get("truncated"):
            raise BoxError("Reports could not be listed completely. The uncertain upload was retained for a safe retry.")
        matches = [item for item in folder.get("entries", []) if item.get("name") == intent.get("name") and item.get("type") == "file"]
        if not matches:
            return None
        if len(matches) != 1:
            raise BoxError("The pending output name is ambiguous. Check Reports before saving again.")
        identifier = box_id(matches[0]["id"])
        info = self.box.file_info(identifier)
        with self.lock:
            protected = {str(source.get("id")) for other in self.runs.values() for source in other.get("sources", [])}
            protected.update(value for key, other in self.runs.items() if key != run["id"] for value in _tracked_outputs(other))
        if identifier in protected or str(info.get("id")) != identifier or info.get("type") != "file" or info.get("item_status") != "active" or info.get("name") != intent.get("name") or str((info.get("parent") or {}).get("id")) != folder_id or str((info.get("owned_by") or {}).get("id")) != str(identity.get("id")) or str(info.get("sha1", "")).lower() not in intent.get("content_sha1", []):
            raise BoxError("A file with the pending output name could not be verified against its intended content, owner and Reports location. It was not adopted or overwritten.")
        return {"id": identifier, "name": info["name"], "url": "https://app.box.com/file/" + identifier}

    def publish(self, run_id):
        with self.operations:
            run, config = self.get(run_id), self.config()
            identity = self.check_identity()
            self.check_versions(run, config)
            digest = _save_key(run)
            label = "Confirmed" if run.get("reviewer_confirmed") else "Draft"
            employee = re.sub(r"[^A-Za-z0-9-]+", "-", run["employee_name"]).strip("-")[:60] or "Employee"
            prefix = f"Annual-Review-{employee}-{run['year']}-{label}-{run_id[:8]}-{digest}"
            files = list(run.get("box_files", []))
            slots = deepcopy(run.get("output_slots", {}))
            pending = deepcopy(run.get("pending_uploads", {}))
            warnings = []
            for kind, format in (("report", "docx"), ("report", "pdf"), ("review_notes", "docx")):
                if any(item.get("save_key") == digest and item.get("kind") == kind and item.get("format") == format for item in _current_box_files({"box_files": files})):
                    continue
                document = _review_document(run) if kind == "review_notes" else run
                try:
                    name = prefix + ("-Review-notes" if kind == "review_notes" else "") + "." + format
                    slot = kind + ":" + format
                    intent_key = digest + ":" + slot
                    recovered = self.recover_upload(run, pending[intent_key], identity, config) if intent_key in pending else None
                    if recovered:
                        uploaded = recovered
                    else:
                        data, _, mime = make_export(document, format)
                    if not recovered and run.get("reuse_output_ids") and slots.get(slot):
                        identifier = box_id(slots[slot])
                        metadata = self.box.file_info(identifier)
                        source_ids = {str(item.get("id")) for item in run.get("sources", [])}
                        with self.lock:
                            shared = any(identifier in _tracked_outputs(other) for key, other in self.runs.items() if key != run_id)
                        if shared or identifier in source_ids or metadata.get("type") != "file" or metadata.get("item_status") != "active" or str((metadata.get("parent") or {}).get("id")) != config["reports_folder_id"] or str((metadata.get("owned_by") or {}).get("id")) != str(identity.get("id")):
                            raise BoxError("The existing output's owner or Reports location changed. Check the linked Box file before saving again.")
                        uploaded = self.box.upload_version(identifier, name, data, mime, etag=str(metadata.get("etag", "")))
                        if str(uploaded.get("id")) != identifier:
                            raise BoxError("Box did not confirm the expected output file. Check its link before saving again.")
                    elif not recovered:
                        previous_intent = deepcopy(pending.get(intent_key))
                        hashes = list((previous_intent or {}).get("content_sha1", []))
                        content_hash = hashlib.sha1(data).hexdigest()
                        if content_hash not in hashes:
                            hashes.append(content_hash)
                        pending[intent_key] = {"save_key": digest, "kind": kind, "format": format, "name": name,
                                               "folder_id": config["reports_folder_id"], "owner_id": str(identity.get("id")), "content_sha1": hashes}
                        self.update(run_id, pending_uploads=pending)
                        try:
                            uploaded = self.box.upload(config["reports_folder_id"], name, data, mime)
                        except BoxError as error:
                            if not error.ambiguous and not previous_intent:
                                pending.pop(intent_key)
                                self.update(run_id, pending_uploads=pending)
                            raise
                    if run.get("reuse_output_ids"):
                        slots[slot] = str(uploaded["id"])
                    files.append({**uploaded, "format": format, "kind": kind, "save_key": digest})
                    pending.pop(intent_key, None)
                    self.update(run_id, box_files=files, output_slots=slots, pending_uploads=pending)
                except Exception as error:
                    warnings.append(_safe_error(error))
            if any(not key.startswith(digest + ":") for key in pending):
                warnings.append("An earlier report upload still has an uncertain outcome. Its private recovery record is retained; draft deletion will verify it before removing history.")
            self.update(run_id, save_warnings=warnings)
            return [_file(item) for item in _current_box_files({"box_files": files}) if item.get("save_key") == digest]

    def execute(self, run_id, config, snapshot, notes):
        try:
            self.update(run_id, status="running", stage="discover", progress=2, message="Reading the configured Box workspace.")
            request = {"folder_id": config["root_folder_id"], "employee_name": config["employee_name"], "year": config["year"], "language": "en", "notes": notes, "max_files": 25}
            result = self.workflow(NativeBoxAdapter(self.box, config["root_folder_id"], snapshot), self.settings, request, lambda **fields: self.update(run_id, **fields))
            self.update(run_id, **{**result, "stage": "save_draft", "progress": 97, "message": "Saving the draft Word, PDF and review notes to Box."})
            self.publish(run_id)
            self.update(run_id, status="completed", stage="completed", progress=100, completed_at=_now(), message="Draft ready for employee review.")
        except Exception as error:
            message = error.detail if isinstance(error, HTTPException) and isinstance(error.detail, str) else _safe_error(error)
            diagnostics = {}
            if isinstance(error, WorkflowError):
                code = error.code if isinstance(error.code, str) else "workflow_error"
                reason = error.validation_reason if isinstance(error.validation_reason, str) else None
                diagnostics["internal_diagnostics"] = {"code": code[:80], "validation_reason": reason[:500] if reason else None}
            self.update(run_id, status="failed", stage="failed", error=message, message=message, **diagnostics)


def create_app(settings=None, box=None, native_config=None, directory=None, workflow=run_workflow, executor=None, improvement=None):
    settings = settings or load_settings()
    service = Workbench(settings, box or BoxClient(settings), native_config, directory, workflow, executor, improvement)

    @asynccontextmanager
    async def lifespan(_app):
        service.load()
        try:
            yield
        finally:
            await asyncio.to_thread(service.shutdown)

    app = FastAPI(title="Arsia | Annual Review Agent", version="0.3.0", lifespan=lifespan)
    app.state.workbench = service
    app.add_middleware(TrustedHostMiddleware, allowed_hosts=["localhost", "127.0.0.1"])
    app.add_middleware(CORSMiddleware, allow_origins=ORIGINS, allow_methods=["GET", "POST", "PUT", "DELETE", "OPTIONS"], allow_headers=["Content-Type"], allow_credentials=False)

    @app.middleware("http")
    async def local_only(request, call_next):
        if request.method in {"POST", "PUT", "DELETE", "PATCH"}:
            origin = request.headers.get("origin")
            if (origin and origin not in ORIGINS) or request.headers.get("sec-fetch-site") == "cross-site":
                return JSONResponse({"detail": "Open the workbench on this computer to make changes."}, status_code=403)
            try:
                size = int(request.headers.get("content-length", "0"))
            except ValueError:
                return JSONResponse({"detail": "Invalid request length."}, status_code=400)
            if size > 125 * 1024 * 1024:
                return JSONResponse({"detail": "Upload request is too large."}, status_code=413)
        response = await call_next(request)
        response.headers["X-Content-Type-Options"] = "nosniff"
        response.headers["Referrer-Policy"] = "no-referrer"
        response.headers["Cache-Control"] = "no-store"
        return response

    @app.exception_handler(BoxError)
    async def box_error(_request, error):
        return JSONResponse({"detail": _safe_error(error)}, status_code=400)

    @app.exception_handler(NativeWorkflowError)
    async def native_error(_request, error):
        return JSONResponse({"detail": _safe_error(error)}, status_code=400)

    @app.exception_handler(Exception)
    async def unexpected_error(_request, _error):
        return JSONResponse({"detail": "The operation could not complete. Check the local configuration and retry."}, status_code=500)

    @app.get("/api/status")
    def status():
        config = service.config()
        categories, reports = [], []
        try:
            identity = service.check_identity()
            snapshot = _snapshot(service.box, config)
            origins = {str(item.get("target_id")): item.get("origin_url") for item in service.import_manifest()}
            categories = [{"name": name, "files": [{**_file({**item, **({"origin_url": origins[item["id"]]} if origins.get(item["id"]) else {})}), "material_id": str(item["id"])} for item in snapshot["files"] if item.get("document_category") == name]} for name in CATEGORIES]
            reports = [_file({**item, "url": "https://app.box.com/file/" + box_id(item["id"])}) for item in service.box.list_folder(config["reports_folder_id"])["entries"] if item.get("type") == "file"]
            box_status = {"connected": True, "identity": {key: identity.get(key) for key in ("name", "type")}, "mode": service.box.mode}
        except BoxError as error:
            box_status = {"connected": False, "error": _safe_error(error)}
        return {"box": box_status, "categories": categories, "employee_name": config["employee_name"], "year": config["year"], "reports": reports, "latest_cloud_reports": reports, "workspace_folder_url": f"https://app.box.com/folder/{config['root_folder_id']}", "reports_folder_url": f"https://app.box.com/folder/{config['reports_folder_id']}", "aws": {"region": settings.aws_region, "model_id": settings.bedrock_model_id, "note": "Model access is checked during generation."}}

    @app.get("/api/box/folders/{folder_id}")
    def browse_box(folder_id: str):
        return service.browse_box(folder_id)

    @app.post("/api/imports")
    def import_box(payload: ImportPayload):
        return service.import_files(payload)

    @app.delete("/api/materials/{file_id}")
    def material_delete(file_id: str):
        return service.delete_material(file_id)

    @app.post("/api/uploads")
    def uploads(category: str = Form(...), files: list[UploadFile] = File(...)):
        with service.operations:
            service.ensure_idle()
            config = service.config()
            service.check_identity()
            if category not in CATEGORIES or not 1 <= len(files) <= 10:
                raise HTTPException(400, "Choose a configured category and 1-10 Word/PDF files.")
            prepared = []
            for file in files:
                name = file.filename or ""
                if Path(name).suffix.lower() not in {".docx", ".pdf"} or not name or any(char in name for char in ("/", "\\", "\r", "\n", "\x00")):
                    raise HTTPException(400, "Only valid Word (.docx) and PDF filenames are accepted.")
                data = file.file.read(MAX_FILE_BYTES + 1)
                if not data or len(data) > MAX_FILE_BYTES:
                    raise HTTPException(400, "Each file must be nonempty and at most 12 MB.")
                prepared.append((name, data, file.content_type or "application/octet-stream"))
            uploaded = []
            for name, data, mime in prepared:
                try:
                    uploaded.append(_file(service.box.upload(config["category_folders"][category], name, data, mime)))
                except BoxError as error:
                    return JSONResponse({"detail": _safe_error(error), "files": uploaded, "partial": bool(uploaded)}, status_code=400)
            return {"files": uploaded, "category": category}

    @app.post("/api/runs", status_code=202)
    def start(payload: RunPayload):
        return service.start(payload.notes)

    @app.get("/api/runs")
    def run_list():
        with service.lock:
            return {"runs": [_public(item) for item in sorted(service.runs.values(), key=lambda item: item.get("created_at", ""), reverse=True)]}

    @app.get("/api/runs/{run_id}")
    def run_get(run_id: str):
        return _public(service.get(run_id))

    @app.delete("/api/runs/{run_id}")
    def run_delete(run_id: str):
        return service.delete_run(run_id)

    def completed(run_id):
        run = service.get(run_id)
        if run.get("status") != "completed":
            raise HTTPException(409, "Wait for generation to complete before reviewing this draft.")
        if run.get("deletion"):
            raise HTTPException(409, "Deletion has started for this draft. Retry deletion to finish before changing or exporting it.")
        return run

    @app.put("/api/runs/{run_id}/report")
    def edit(run_id: str, payload: ReportPayload):
        with service.operations:
            run = completed(run_id)
            service.check_report_version(run, payload.expected_report_hash)
            if payload.report_markdown == run.get("report_markdown"):
                return _public(run)
            service.update(run_id, report_markdown=payload.report_markdown, reviewer_confirmed=False, confirmed_at=None, user_edited=True, message="Employee edited the draft. Confirmation has been reset.")
            return _public(service.get(run_id))

    @app.post("/api/runs/{run_id}/confirm")
    def confirm(run_id: str, payload: VersionPayload):
        with service.operations:
            run = completed(run_id)
            service.check_report_version(run, payload.expected_report_hash)
            service.check_versions(run, service.config())
            service.update(run_id, reviewer_confirmed=True, confirmed_at=_now(), message="The employee explicitly confirmed this version.")
            return _public(service.get(run_id))

    @app.get("/api/runs/{run_id}/exports/{format}")
    def export(run_id: str, format: Literal["docx", "pdf"], prepare: bool = False):
        with service.operations:
            run = completed(run_id)
            service.check_versions(run, service.config())
            data, name, mime = make_export(run, format)
            if prepare:
                return {"download_url": f"/api/runs/{run_id}/exports/{format}", "filename": name, "content_type": mime}
            return Response(data, media_type=mime, headers={"Content-Disposition": f'attachment; filename="{name}"'})

    @app.post("/api/runs/{run_id}/save")
    def save(run_id: str, payload: SavePayload = SavePayload()):
        run = completed(run_id)
        if payload.reviewer_confirmed and not run.get("reviewer_confirmed"):
            raise HTTPException(409, "Confirm this report explicitly before saving it as confirmed.")
        files = service.publish(run_id)
        return {"files": files, "run": _public(service.get(run_id))}

    @app.post("/api/runs/{run_id}/answers", status_code=202)
    @app.post("/api/runs/{run_id}/improve", status_code=202)
    def answers(run_id: str, payload: AnswerPayload):
        return service.improve(run_id, payload.answers, payload.target_score, payload.expected_report_hash)

    @app.get("/{path:path}", include_in_schema=False)
    def frontend(path: str):
        if path == "api" or path.startswith("api/"):
            raise HTTPException(404, "API route not found.")
        directory = (ROOT / "frontend" / "dist").resolve()
        candidate = (directory / path).resolve()
        if not candidate.is_relative_to(directory):
            raise HTTPException(404, "Page not found.")
        if candidate.is_file():
            return FileResponse(candidate)
        index = directory / "index.html"
        if not index.is_file() or Path(path).suffix:
            raise HTTPException(404, "Build the local frontend first with the start script.")
        return FileResponse(index)

    return app


app = create_app()
