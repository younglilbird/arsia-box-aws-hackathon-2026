#!/usr/bin/env python3
"""Build the nine real Office/PDF demo inputs from versioned source text.

Run from any directory: python scripts/build_demo_documents.py
Only the resulting files/ directory is uploaded to Box by the demo workflow.
The Markdown files remain maintainable authoring sources, not demo inputs.
"""

from __future__ import annotations

import re
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "demo-materials"
OUTPUT = SOURCE / "files"
SPEC = [
    ("01-employee-and-reporting-context", "pdf", "Alex Chen review preparation context"),
    ("02-atlas-project-design", "docx", "Atlas guided workspace onboarding design"),
    ("03-alex-atlas-work-log", "docx", "Atlas measurement work by Alex Chen"),
    ("04-atlas-pilot-results", "pdf", "Atlas pilot measurement readout"),
    ("05-atlas-stale-status-snapshot", "pdf", "Atlas status snapshot"),
    ("06-beacon-design-and-contributions", "docx", "Beacon support knowledge search"),
    ("07-beacon-pilot-evaluation", "pdf", "Beacon pilot evaluation"),
    ("08-annual-product-plans", "docx", "Proposed annual product directions"),
    ("09-peer-feedback", "pdf", "Peer feedback for Alex Chen"),
]


def load_source(path: Path):
    text = path.read_text(encoding="utf-8")
    metadata = {}
    if text.startswith("---\n"):
        _, front, text = text.split("---", 2)
        metadata = dict(line.split(":", 1) for line in front.splitlines() if ":" in line)
        metadata = {key.strip(): value.strip() for key, value in metadata.items()}
    blocks = []
    pending = []
    lines = text.splitlines()
    index = 0

    def flush():
        if pending:
            blocks.append(("paragraph", " ".join(pending)))
            pending.clear()

    while index < len(lines):
        line = lines[index].strip()
        index += 1
        if not line:
            flush()
        elif line.startswith("# "):
            flush()  # the curated document title replaces the source title
        elif line.startswith("## "):
            flush()
            blocks.append(("heading", re.sub(r"[^\w\s]", "", line[3:])))
        elif line.startswith("| "):
            flush()
            rows = []
            while True:
                cells = [cell.strip() for cell in line.strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                    rows.append(cells)
                if index >= len(lines) or not lines[index].strip().startswith("|"):
                    break
                line = lines[index].strip()
                index += 1
            blocks.append(("table", rows))
        elif line.startswith("- "):
            flush()
            blocks.append(("bullet", line[2:]))
        else:
            pending.append(line)
    flush()
    return metadata, blocks


def metadata_lines(metadata: dict):
    fields = [
        ("document_date", "Date"),
        ("author", "Author"), ("subject", "Subject"), ("project", "Project"),
        ("status", "Status"),
    ]
    return [f"{label}: {metadata[key].replace('_', ' ')}" for key, label in fields if metadata.get(key)]


def _clean_title_style(style):
    style.font.color.rgb = RGBColor(0, 0, 0)
    properties = style.element.find(qn("w:pPr"))
    if properties is not None:
        border = properties.find(qn("w:pBdr"))
        if border is not None:
            properties.remove(border)


def build_docx(path: Path, title: str, metadata: dict, blocks: list):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.75)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(7)
    for name, size in [("Title", 23), ("Heading 1", 13), ("Heading 2", 12)]:
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = name != "Title"
        _clean_title_style(style)
        style.paragraph_format.space_before = Pt(11 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    document.add_paragraph(title, "Title")
    for item in metadata_lines(metadata):
        paragraph = document.add_paragraph(item)
        paragraph.paragraph_format.space_after = Pt(2)
        for run in paragraph.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string("4A4A4A")
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    for kind, value in blocks:
        if kind == "heading":
            document.add_paragraph(value, "Heading 1")
        elif kind == "bullet":
            document.add_paragraph(value, "List Bullet")
        elif kind == "table":
            table = document.add_table(rows=0, cols=len(value[0]))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            widths = [Inches(3.5), Inches(1.75), Inches(1.75)]
            for row_index, row_values in enumerate(value):
                cells = table.add_row().cells
                for col_index, cell_text in enumerate(row_values):
                    cell = cells[col_index]
                    cell.width = widths[col_index]
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    cell.text = cell_text
                    properties = cell._tc.get_or_add_tcPr()
                    borders = OxmlElement("w:tcBorders")
                    for side in ("top", "left", "bottom", "right"):
                        edge = OxmlElement(f"w:{side}")
                        edge.set(qn("w:val"), "single")
                        edge.set(qn("w:sz"), "4")
                        edge.set(qn("w:color"), "D9D9D9")
                        borders.append(edge)
                    properties.append(borders)
                    margins = OxmlElement("w:tcMar")
                    for side in ("top", "left", "bottom", "right"):
                        margin = OxmlElement(f"w:{side}")
                        margin.set(qn("w:w"), "100")
                        margin.set(qn("w:type"), "dxa")
                        margins.append(margin)
                    properties.append(margins)
                    shading = OxmlElement("w:shd")
                    shading.set(qn("w:fill"), "DCE6F1" if row_index == 0 else "F7F9FB" if row_index % 2 == 0 else "FFFFFF")
                    properties.append(shading)
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(0)
                        for run in paragraph.runs:
                            run.bold = row_index == 0
                            run.font.size = Pt(10)
            document.add_paragraph()
        else:
            document.add_paragraph(value)
    document.core_properties.title = title
    document.core_properties.subject = "Synthetic Northstar Cloud demonstration document"
    document.core_properties.author = metadata.get("author", "")
    document.core_properties.keywords = "synthetic, demonstration, fictional data"
    document.save(path)


def build_pdf(path: Path, title: str, metadata: dict, blocks: list):
    normal = ParagraphStyle("Body", fontName="Helvetica", fontSize=11, leading=15.5, textColor=colors.black, spaceAfter=8)
    title_style = ParagraphStyle("Title", parent=normal, fontName="Helvetica-Bold", fontSize=22, leading=27, spaceAfter=12)
    heading = ParagraphStyle("Heading", parent=normal, fontName="Helvetica-Bold", fontSize=13, leading=17, spaceBefore=10, spaceAfter=6, keepWithNext=True)
    small = ParagraphStyle("Metadata", parent=normal, fontSize=9, leading=12, spaceAfter=2, textColor=colors.HexColor("#4A4A4A"))
    cell = ParagraphStyle("Cell", parent=normal, fontSize=10, leading=13, spaceAfter=0)
    cell_heading = ParagraphStyle("CellHeading", parent=cell, fontName="Helvetica-Bold")
    document = SimpleDocTemplate(str(path), pagesize=letter, leftMargin=54, rightMargin=54, topMargin=47, bottomMargin=47, title=title, author=metadata.get("author", ""), subject="Synthetic Northstar Cloud demonstration document")
    story = [Paragraph(escape(title), title_style)]
    story.extend(Paragraph(escape(item), small) for item in metadata_lines(metadata))
    story.append(Spacer(1, 14))
    for kind, value in blocks:
        if kind == "heading":
            story.append(Paragraph(escape(value), heading))
        elif kind == "bullet":
            story.append(Paragraph("- " + escape(value), normal))
        elif kind == "table":
            data = [[Paragraph(escape(text), cell_heading if index == 0 else cell) for text in row] for index, row in enumerate(value)]
            table = Table(data, colWidths=[252, 126, 126], repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6F1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F9FB")]),
                ("GRID", (0, 0), (-1, -1), .5, colors.HexColor("#D9D9D9")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ]))
            story.extend([Spacer(1, 3), table, Spacer(1, 12)])
        else:
            story.append(Paragraph(escape(value), normal))
    document.build(story)


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    expected = {f"{stem}.{format}" for stem, format, _ in SPEC}
    for old in OUTPUT.iterdir():
        if old.is_file() and old.suffix in {".pdf", ".docx"} and old.name not in expected:
            old.unlink()
    for stem, format, title in SPEC:
        metadata, blocks = load_source(SOURCE / f"{stem}.md")
        path = OUTPUT / f"{stem}.{format}"
        (build_docx if format == "docx" else build_pdf)(path, title, metadata, blocks)
        print(f"{path.relative_to(ROOT)} ({path.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
